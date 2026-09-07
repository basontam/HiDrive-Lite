#!/usr/bin/env python3
"""Offline Excel importer for the personal media-resource library
(T2.5; docs/architecture.md §4.8, §5).

    .venv/bin/python scripts/import_media_library.py
        --source-dir /path/to/your/workbooks
        --output .local-index                      # bundle dir (gitignored)
        [--write]                                   # default: dry-run
        [--report build/test-artifacts/library-import-report.json]
        [--require-hashes]
        [--limit N]

Reads the three source workbooks (see ``SOURCES``), aggregates everything in
memory following the §4.8 dedup/merge order, then -- only with ``--write``
-- writes a fresh ``library-bundle.sqlite`` through ``library_store`` and
builds the search index via ``library_search.build_index``.  A desensitised
JSON report (counts and hashes only, never a link or access-code value) is
always written to ``--report`` when given.

This is the only module in the project allowed to import ``openpyxl``,
``pypinyin`` and ``zhconv`` -- all three are dev-only dependencies
(``requirements-dev.txt``) and are imported lazily, inside the functions
that actually need them, so a caller that only wants ``read_sheet``/
``aggregate`` (e.g. a test) never pays for pypinyin/zhconv import cost, and
so it stays obvious this is the one place they are allowed to appear.
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import os
import re
import sys
import time
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from library_normalize import (  # noqa: E402
    NORMALIZE_VERSION,
    EditionInfo,
    TitleInfo,
    canonical_hash,
    edition_fingerprint,
    infer_media_type,
    media_identity,
    parse_edition,
    parse_link,
    parse_title,
    public_id,
    search_key,
)
from library_store import (  # noqa: E402
    _use_temp_dir,
    GroupRecord,
    LibraryStore,
    LinkRecord,
    MediaRecord,
    ProvenanceRecord,
)

HEADER = [
    "类型", "记录ID", "Slug/追更包Slug", "用户UID", "官组标记", "标题",
    "媒体标题", "链接", "访问码", "备注", "创建时间", "删除时间",
]

# (filename, sheet, kind) -- processed in this exact order (§4.8 rule 1):
# main table first, then the two supplements.  "kind" is also the key used
# in aggregate()'s rows_by_source dict.
SOURCES = [
    ("HDHive official-group-backup.xlsx", "影视分享", "main"),
    ("影巢_影视分享_整理.xlsx", "影视分享", "supplement_video"),
    ("影巢_ED2K_整理.xlsx", "磁力ED2K", "supplement_ed2k"),
]

# The official workbook's 磁力ED2K sheet is a filtered view of 分享明细, not
# imported directly; it is only used to cross-check that every (Slug,
# canonical link) pair it contains already appeared while processing the
# main table.
CROSS_CHECK = ("HDHive official-group-backup.xlsx", "磁力ED2K")

_DEFAULT_REPORT_PATH = Path("build/test-artifacts/library-import-report.json")


class ImportRowError(Exception):
    """Raised for a row/column that failed to parse.

    Carries only ``sheet``, ``row_number`` and ``column`` -- never the cell
    value itself, so this exception (and anything that logs ``str()`` of
    it) can never leak a link, access code or title.
    """

    def __init__(self, sheet: str, row_number: int, column: str) -> None:
        self.sheet = sheet
        self.row_number = row_number
        self.column = column
        super().__init__(f"import error: sheet={sheet!r} row={row_number} column={column!r}")


class SourceFileMissingError(Exception):
    """Named only so ``run()``'s failure report can put a real class name
    in its ``"error"`` field -- never actually raised; the missing-file
    check already prints to stderr and returns before any exception
    handling would apply."""


class SourceHashMismatchError(Exception):
    """Named only so ``run()``'s failure report can put a real class name
    in its ``"error"`` field when ``--require-hashes`` is set and a source
    workbook's sha256 does not match the README's reference table -- never
    actually raised, see :class:`SourceFileMissingError`."""


# ---------------------------------------------------------------------------
# Aggregate -- everything the importer builds up in memory before writing.
# ---------------------------------------------------------------------------


@dataclass
class Aggregate:
    # media_identity -> {media_type, title_zh, title_original, aliases(set),
    #                     year, search_key, needs_review}
    media: dict = field(default_factory=dict)
    # (media_identity, edition_fingerprint) -> group fields
    groups: dict = field(default_factory=dict)
    # (provider, canonical_hash) -> link fields
    links: dict = field(default_factory=dict)
    # flat list of {link_key, source_file, sheet, row_number, record_id,
    #                slug, owner_uid, owner_tag}
    provenance: list = field(default_factory=list)
    # kind -> {rows_total, rows_video, rows_skipped_non_video,
    #           rows_invalid_link, rows_deleted}
    sources: dict = field(default_factory=dict)
    providers: Counter = field(default_factory=Counter)
    media_type: Counter = field(default_factory=Counter)
    year_missing: int = 0
    dedup: dict = field(default_factory=dict)
    conflicts: dict = field(default_factory=dict)
    needs_review: dict = field(default_factory=dict)
    totals: dict = field(default_factory=dict)
    elapsed_seconds: float = 0.0


# ---------------------------------------------------------------------------
# read_sheet
# ---------------------------------------------------------------------------


def read_sheet(path: Path, sheet: str) -> tuple:
    """Locate the header row (first row containing a "记录ID" cell), verify
    all 12 §1.2 column names are present, and return
    ``(header_row_number, rows)`` where each row is a dict keyed by column
    name (plus an internal ``"_row_number"`` key -- the row's real position
    in the sheet, used for provenance and for ``ImportRowError``).  Rows
    whose "记录ID" cell is empty are skipped entirely.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ImportRowError(sheet, 0, "sheet")
        ws = wb[sheet]

        header_row_number = None
        header_values: list = []
        for row in ws.iter_rows(min_row=1, max_row=20):
            # Minor: tolerate a trailing space or a leading BOM ("﻿",
            # which Excel/CSV-round-tripped workbooks sometimes leave on the
            # very first header cell) in a header cell -- data rows are read
            # as-is, only header detection/column matching is normalised.
            values = [
                cell.value.replace("﻿", "").strip()
                if isinstance(cell.value, str) else cell.value
                for cell in row
            ]
            if any(v == "记录ID" for v in values):
                header_row_number = row[0].row
                header_values = values
                break
        if header_row_number is None:
            raise ImportRowError(sheet, 0, "记录ID")

        col_index: dict = {}
        for idx, value in enumerate(header_values):
            if value in HEADER and value not in col_index:
                col_index[value] = idx
        missing = [name for name in HEADER if name not in col_index]
        if missing:
            raise ImportRowError(sheet, header_row_number, missing[0])

        rows: list = []
        for row in ws.iter_rows(min_row=header_row_number + 1):
            def cell_value(name, _row=row):
                idx = col_index[name]
                return _row[idx].value if idx < len(_row) else None

            record_id = cell_value("记录ID")
            if record_id is None or (isinstance(record_id, str) and not record_id.strip()):
                continue
            row_dict = {name: cell_value(name) for name in HEADER}
            row_dict["_row_number"] = row[0].row
            rows.append(row_dict)
        return header_row_number, rows
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# aggregate
# ---------------------------------------------------------------------------


def _parse_source_timestamp_ex(value) -> tuple:
    """``(epoch_seconds_or_None, is_bad)`` for a source-workbook 创建时间/
    删除时间 cell.  ``is_bad`` is True only for a non-empty, non-"NULL"
    string that matches none of the accepted formats -- real-data hotfix:
    the 影巢 supplement workbooks hold 创建时间/删除时间 as strings (e.g.
    ``"2025-05-08 15:35:13"``), not ``datetime`` objects like the main
    workbook, in ``"YYYY-MM-DD HH:MM:SS"``/``"YYYY-MM-DDTHH:MM:SS"``/
    ``"YYYY-MM-DD"`` (all via ``datetime.fromisoformat``, which already
    tolerates case and any single separator character) or
    ``"YYYY/MM/DD HH:MM:SS"``.  Never raises -- aggregate() uses ``is_bad``
    to fill ``sources[].rows_bad_timestamp`` and the needs_review
    ``"timestamp_unparsed"`` reason instead of failing the whole import, as
    it used to (see ``ImportRowError``).  ``parse_source_timestamp`` below
    is the public value-only wrapper other callers should use."""
    if value is None or value == "" or isinstance(value, bool):
        return None, False
    if isinstance(value, (int, float)):
        return int(value), False
    if isinstance(value, datetime):
        return int(value.timestamp()), False
    if isinstance(value, str):
        v = value.strip()
        if not v or v.upper() == "NULL":
            return None, False
        try:
            return int(datetime.fromisoformat(v).timestamp()), False
        except ValueError:
            pass
        try:
            return int(datetime.strptime(v, "%Y/%m/%d %H:%M:%S").timestamp()), False
        except ValueError:
            return None, True
    return None, True


def parse_source_timestamp(value) -> int | None:
    """Parse a source-workbook 创建时间/删除时间 cell into epoch seconds, or
    ``None`` for an empty cell -- see ``_parse_source_timestamp_ex`` for the
    accepted formats.  Never raises: a non-empty string that matches none of
    them also returns ``None`` (callers that need to distinguish that case
    from a genuinely empty cell use ``_parse_source_timestamp_ex``)."""
    return _parse_source_timestamp_ex(value)[0]


def _parse_int(value, *, sheet: str, row_number: int, column: str):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ImportRowError(sheet, row_number, column)
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        v = value.strip()
        if not v:
            return None
        try:
            return int(v)
        except ValueError:
            raise ImportRowError(sheet, row_number, column) from None
    raise ImportRowError(sheet, row_number, column)


def _cell_str(value):
    """Coerce a cell that should hold text into a stripped ``str``, or
    ``None`` for a blank cell -- real-data hotfix (I7): a handful of rows in
    the source workbooks hold a text column (标题/媒体标题/访问码/备注/
    链接/类型) as a numeric cell (openpyxl hands back an ``int``/``float``
    for a cell Excel auto-typed as a number, e.g. an all-digit 访问码)
    instead of a string -- this used to raise deep inside parse_title/
    parse_edition/parse_link (``AttributeError: 'int' object has no
    attribute 'strip'`` and friends) and abort the whole run.  A whole-number
    ``float`` renders without the trailing ``".0"`` (``500001.0`` ->
    ``"500001"``), matching how the same digits would read as a string
    cell."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


# Real-data hotfix: 12 rows in the main workbook are shifted one column
# right starting at 链接 -- 链接 is empty, 访问码 holds the URL, 备注 holds
# the (optional) access code, 创建时间 holds the remark text and 删除时间
# holds the real creation datetime.  See _is_row_shifted/_repair_shifted_row.
_SHIFTED_LINK_SCHEMES = ("http://", "https://", "ed2k://", "magnet:")
_SHIFTED_ACCESS_CODE_RE = re.compile(r"^[A-Za-z0-9]{1,8}$")


def _is_row_shifted(row: dict) -> bool:
    """Detected structurally: an empty -- or, once stripped, whitespace-only
    (I10) -- 链接 cell paired with a 访问码 cell that itself looks like a
    link (case-insensitive scheme check)."""
    link_cell = row.get("链接")
    if isinstance(link_cell, str):
        link_cell = link_cell.strip()
    if link_cell:
        return False
    access_code_cell = row.get("访问码")
    if not isinstance(access_code_cell, str):
        return False
    return access_code_cell.strip().lower().startswith(_SHIFTED_LINK_SCHEMES)


def _repair_shifted_row(row: dict) -> dict:
    """Un-shift a row flagged by ``_is_row_shifted``: link <- 访问码, access
    code <- 备注 (only if it looks like a real access code, else None),
    remark <- 创建时间 (the real remark text), created <- 删除时间 (the real
    creation timestamp).  The real deletion time was never captured by the
    shift, so it is always lost here -- treated as None.  Returns a new
    dict; never mutates ``row``."""
    repaired = dict(row)
    repaired["链接"] = row.get("访问码")
    remark_as_code = row.get("备注")
    if isinstance(remark_as_code, str) and _SHIFTED_ACCESS_CODE_RE.match(remark_as_code.strip()):
        repaired["访问码"] = remark_as_code.strip()
    else:
        repaired["访问码"] = None
    repaired["备注"] = row.get("创建时间")
    repaired["创建时间"] = row.get("删除时间")
    repaired["删除时间"] = None
    return repaired


# Real-data hotfix: 161 rows pack 2-142 ED2K/magnet/http(s) links into one
# 链接 cell with no separator openpyxl can rely on (ED2K file names can
# themselves contain spaces) -- split on a lookahead for each recognised
# scheme so every link is addressable. See _split_link_fragments.
_LINK_FRAGMENT_SPLIT_RE = re.compile(r"(?=ed2k://|magnet:|https?://)", re.IGNORECASE)


def _split_link_fragments(link_cell: str) -> list:
    """Split a (non-blank) 链接 cell that may pack multiple links into one
    string on a lookahead for ``ed2k://``/``magnet:``/``http(s)://``.  A
    cell with none of these schemes at all (free text -- the existing
    invalid-link case) yields itself unchanged as the sole "fragment", so
    the overwhelmingly common single-link cell is unaffected."""
    parts = [p.strip() for p in _LINK_FRAGMENT_SPLIT_RE.split(link_cell) if p.strip()]
    return parts or [link_cell]


def _title_alias_conflict(title_cell: str, media_title_cell: str) -> bool:
    """§4.8 rule 6: 媒体标题 and 标题去年份's search_key neither contain the
    other.  ``parse_title(title_cell, "")`` isolates "标题去年份" (year
    stripped, alias-split untouched) via the same rule ``parse_title`` uses
    internally, without duplicating its year-stripping regex here."""
    media_title_cell = (media_title_cell or "").strip()
    if not media_title_cell:
        return False
    title_no_year = parse_title(title_cell, "").title_zh
    media_key = search_key(media_title_cell)
    title_key = search_key(title_no_year)
    if not media_key or not title_key:
        return False
    return media_key not in title_key and title_key not in media_key


# ---------------------------------------------------------------------------
# edition_display_title -- resource_group.display_title (§3 DDL).
#
# Kept here rather than in library_normalize.py: this task's scope is
# scripts/import_media_library.py + its test file only, and
# library_normalize.py belongs to a different task/worktree (t1-normalize).
# edition_display_title is a pure function of EditionInfo with no
# openpyxl/pypinyin/zhconv dependency, so there is no cost to keeping it
# here instead.
# ---------------------------------------------------------------------------

_QUALITY_DISPLAY = {"2160p": "4K", "1080p": "1080p", "720p": "720p"}
_SOURCE_TYPE_DISPLAY = {
    "remux": "REMUX", "bluray": "BluRay", "bdrip": "BDRip",
    "webdl": "WEB-DL", "hdtv": "HDTV",
}
_HDR_DISPLAY = {
    "dv_hdr": "DV/HDR", "dv": "DV", "hdr10plus": "HDR10+",
    "hdr10": "HDR", "hlg": "HLG", "sdr": "SDR",
}
_CODEC_DISPLAY = {"hevc": "HEVC", "avc": "AVC", "av1": "AV1"}
# Two remark phrases can fold to the same internal subtitle code (§4.3);
# each code's display form here is its first-listed source phrase.
_SUBTITLE_DISPLAY = {
    "embedded:zh-hans+zh-hant": "内封简繁",
    "embedded:zh-hans": "内封简中",
    "embedded:zh-hant": "内封繁中",
    "embedded:zh-hans+en": "内封简英",
    "external:zh-hans": "外挂简中",
    "effects": "特效字幕",
    "dual:zh-hans+en": "简英双语",
    "embedded:ai:zh-hans+zh-hant": "内封AI简繁",
    "zh": "中字",
    "zh-hans+zh-hant": "简/繁",
}


def _edition_season_episode_part(edition: EditionInfo) -> str | None:
    """The season/episode-range fragment of ``edition_display_title``, e.g.
    ``"S01"``, ``"S01-S04"``, ``"S01E01-E06"``, ``"第3集"`` or ``"全季"``
    (the last only when nothing more specific than "complete" is known)."""
    if edition.season_from is not None:
        if edition.season_to is not None and edition.season_to != edition.season_from:
            return f"S{edition.season_from:02d}-S{edition.season_to:02d}"
        season = f"S{edition.season_from:02d}"
        if edition.episode_from is None:
            return season
        if edition.episode_to is not None and edition.episode_to != edition.episode_from:
            return f"{season}E{edition.episode_from:02d}-E{edition.episode_to:02d}"
        return f"{season}E{edition.episode_from:02d}"
    if edition.episode_from is not None:
        if edition.episode_to is not None and edition.episode_to != edition.episode_from:
            return f"E{edition.episode_from:02d}-E{edition.episode_to:02d}"
        return f"第{edition.episode_from}集"
    if edition.complete_season:
        return "全季"
    return None


def edition_display_title(edition: EditionInfo) -> str:
    """§3 DDL: ``resource_group.display_title`` -- the edition summary
    (season/episode, quality, source, HDR, codec, subtitle), never the
    media title.  Parts are joined with " · "; ``"未标注版本"`` when
    nothing at all is recognised (matches ``edition_fingerprint``'s own
    ``"unspecified"`` case)."""
    parts: list = []
    season_part = _edition_season_episode_part(edition)
    if season_part:
        parts.append(season_part)
    if edition.quality in _QUALITY_DISPLAY:
        parts.append(_QUALITY_DISPLAY[edition.quality])
    if edition.source_type in _SOURCE_TYPE_DISPLAY:
        parts.append(_SOURCE_TYPE_DISPLAY[edition.source_type])
    if edition.hdr in _HDR_DISPLAY:
        parts.append(_HDR_DISPLAY[edition.hdr])
    if edition.video_codec in _CODEC_DISPLAY:
        parts.append(_CODEC_DISPLAY[edition.video_codec])
    if edition.subtitle:
        parts.append("+".join(sorted(_SUBTITLE_DISPLAY.get(s, s) for s in edition.subtitle)))
    if not parts:
        return "未标注版本"
    return " · ".join(parts)


def _ensure_media(agg: Aggregate, identity: str, info, media_type: str) -> dict:
    entry = agg.media.get(identity)
    if entry is None:
        entry = {
            "media_identity": identity,
            "media_type": media_type,
            "title_zh": info.title_zh,
            "title_original": info.original_hint,
            "aliases": set(),
            "year": info.year,
            "search_key": search_key(info.title_zh),
            "needs_review": False,
        }
        agg.media[identity] = entry
    entry["aliases"].update(a for a in info.aliases if a)
    if info.original_hint and not entry["title_original"]:
        entry["title_original"] = info.original_hint
    return entry


def _ensure_group(agg: Aggregate, group_key: tuple, edition) -> dict:
    entry = agg.groups.get(group_key)
    if entry is None:
        entry = {
            "media_identity": group_key[0],
            "edition_fingerprint": group_key[1],
            "display_title": edition_display_title(edition),
            "season_from": edition.season_from,
            "season_to": edition.season_to,
            "episode_from": edition.episode_from,
            "episode_to": edition.episode_to,
            "complete_season": bool(edition.complete_season),
            "quality": edition.quality,
            "source_type": edition.source_type,
            "hdr": edition.hdr,
            "video_codec": edition.video_codec,
            "audio": edition.audio,
            "subtitle": edition.subtitle,
            "tags": edition.tags,
            "needs_review": False,
            "review_reasons": set(),
        }
        agg.groups[group_key] = entry
    return entry


def _mark_review(agg: Aggregate, group_key: tuple, identity: str, reasons: set) -> None:
    group = agg.groups.get(group_key)
    if group is not None:
        group["needs_review"] = True
        group["review_reasons"].update(reasons)
    media = agg.media.get(identity)
    if media is not None:
        media["needs_review"] = True


def _main_row_sort_key(row: dict) -> tuple:
    """§4.8 rule 1: 主表 影视分享（按 创建时间、记录ID 升序）-- sort the
    main-table rows by 创建时间 ascending (None last), then 记录ID
    ascending, before aggregation.  Only "main" rows are sorted; the two
    supplement sources stay in sheet order per the same rule.  A 创建时间
    that fails to parse (garbage string) sorts as if it were None (last);
    the same cell is re-read via ``parse_source_timestamp`` during
    aggregation, where it is counted/flagged rather than raising."""
    created_ts = parse_source_timestamp(row.get("创建时间"))
    try:
        record_id_sort = int(row.get("记录ID"))
    except (TypeError, ValueError):
        record_id_sort = 0
    if created_ts is None:
        return (1, 0, record_id_sort)
    return (0, created_ts, record_id_sort)


def aggregate(rows_by_source: dict) -> Aggregate:
    """Aggregate all rows in memory, following §4.8's dedup/merge order.

    ``rows_by_source`` is keyed by SOURCES' "kind" ("main",
    "supplement_video", "supplement_ed2k") plus an optional "cross_check"
    key holding the official workbook's 磁力ED2K sheet rows (used only for
    the §4.8 rule-1 cross-check, never aggregated into media/links).
    """
    agg = Aggregate()
    agg.dedup = {
        "within_main_duplicate_links": 0,
        "supplement_video_matched": 0, "supplement_video_added": 0,
        "supplement_ed2k_matched": 0, "supplement_ed2k_added": 0,
        "cross_check_ed2k_sheet_subset": True,
    }
    agg.conflicts = {
        "link_shared_across_groups": 0, "access_code_conflict": 0,
        "access_code_divergence": 0,
    }
    agg.needs_review = {
        "year_missing": 0, "title_alias_conflict": 0, "edition_unparsed": 0,
        "link_shared_across_groups": 0, "access_code_conflict": 0,
        "timestamp_unparsed": 0, "row_shifted": 0,
    }

    main_pairs: set = set()

    for filename, sheet, kind in SOURCES:
        rows = rows_by_source.get(kind, [])
        if kind == "main":
            rows = sorted(rows, key=_main_row_sort_key)
        stat = {
            "rows_total": len(rows), "rows_video": 0, "rows_skipped_non_video": 0,
            "rows_type_missing": 0, "rows_invalid_link": 0, "rows_missing_link": 0,
            "rows_deleted": 0, "rows_bad_timestamp": 0, "rows_shifted": 0,
            "rows_missing_record_id": 0, "rows_multi_link": 0,
            "links_from_multi_link_cells": 0, "rows_error": 0,
        }
        for row in rows:
            row_number = row.get("_row_number", 0)
            # I12: 类型 compared after str(...).strip() -- a None/blank cell
            # is counted separately (rows_type_missing) from a real non-影视
            # type (rows_skipped_non_video).
            type_cell = _cell_str(row.get("类型"))
            if type_cell != "影视":
                if type_cell:
                    stat["rows_skipped_non_video"] += 1
                else:
                    stat["rows_type_missing"] += 1
                continue
            stat["rows_video"] += 1

            # I7: everything below can raise ImportRowError (a cell
            # parse_title/parse_edition/parse_link/_parse_int can't handle)
            # -- contained per row so one bad row never aborts the whole
            # source: counted in rows_error, the row is skipped, the run
            # continues.
            try:
                shifted = _is_row_shifted(row)
                if shifted:
                    row = _repair_shifted_row(row)
                    stat["rows_shifted"] += 1

                # I7: a numeric/float cell in a text column (openpyxl hands
                # back an int/float for a cell Excel auto-typed as a number,
                # e.g. an all-digit 访问码) used to raise deep inside
                # parse_title/parse_edition/parse_link -- _cell_str coerces
                # every text column to a stripped str (or None) first.
                title_cell = _cell_str(row.get("标题")) or ""
                media_title_cell = _cell_str(row.get("媒体标题")) or ""
                remark_cell = _cell_str(row.get("备注")) or ""
                link_cell = _cell_str(row.get("链接")) or ""
                code_cell = _cell_str(row.get("访问码"))

                # I9/I10: a 链接 cell that is blank -- or, once stripped,
                # whitespace-only -- after shift-repair must never reach
                # parse_link: every such row would otherwise canonicalise to
                # the same "unknown:<sha256 of \"\">" link and silently merge
                # into one (or conflict) across unrelated rows.
                if not link_cell:
                    stat["rows_missing_link"] += 1
                    continue

                # Any of these can raise on real-world cell content the
                # parser doesn't expect (e.g. an ED2K file name urlparse
                # chokes on); re-raise as ImportRowError(sheet, row_number,
                # column) with the original exception suppressed (`from
                # None`) so no traceback ever carries a cell value
                # (title/link/access-code/remark).
                try:
                    info = parse_title(title_cell, media_title_cell)
                except Exception:
                    raise ImportRowError(sheet, row_number, "标题") from None
                try:
                    edition = parse_edition(remark_cell, title_cell)
                except Exception:
                    raise ImportRowError(sheet, row_number, "备注") from None

                m_type = infer_media_type(edition, f"{title_cell} {media_title_cell} {remark_cell}")
                identity = media_identity(info, m_type, None)
                fingerprint = edition_fingerprint(edition)
                group_key = (identity, fingerprint)

                created_at, created_bad = _parse_source_timestamp_ex(row.get("创建时间"))
                deleted_at, deleted_bad = _parse_source_timestamp_ex(row.get("删除时间"))
                # Minor: rows_deleted counts a 删除时间 that actually
                # parsed, not the cell's raw truthiness -- a garbage string
                # is truthy but produces no usable timestamp.
                if deleted_at is not None:
                    stat["rows_deleted"] += 1
                if created_bad or deleted_bad:
                    stat["rows_bad_timestamp"] += 1

                reasons_base: set = set()
                if info.year is None:
                    reasons_base.add("year_missing")
                if _title_alias_conflict(title_cell, media_title_cell):
                    reasons_base.add("title_alias_conflict")
                if edition.unparsed:
                    reasons_base.add("edition_unparsed")
                if created_bad or deleted_bad:
                    reasons_base.add("timestamp_unparsed")
                if shifted:
                    reasons_base.add("row_shifted")

                record_id = _parse_int(row.get("记录ID"), sheet=sheet, row_number=row_number, column="记录ID")
                if record_id is None:
                    stat["rows_missing_record_id"] += 1
                owner_uid = _parse_int(row.get("用户UID"), sheet=sheet, row_number=row_number, column="用户UID")
                slug_raw = row.get("Slug/追更包Slug")
                slug = str(slug_raw).strip() if slug_raw not in (None, "") else None
                owner_tag = row.get("官组标记") or None

                # I11: a 链接 cell can pack several links back to back (real
                # data: 2-142 ED2K links in one cell) -- split on a
                # lookahead for each recognised scheme so every one is
                # addressable; for the overwhelmingly common single-link
                # cell this is a one-element list, so behaviour there is
                # unchanged.
                fragments = _split_link_fragments(link_cell)
                multi = len(fragments) > 1
                if multi:
                    stat["rows_multi_link"] += 1

                for fragment in fragments:
                    try:
                        link = parse_link(fragment, code_cell)
                    except Exception:
                        raise ImportRowError(sheet, row_number, "链接") from None

                    if link.provider == "unknown":
                        stat["rows_invalid_link"] += 1
                    if multi:
                        stat["links_from_multi_link_cells"] += 1

                    link_key = (link.provider, canonical_hash(link.canonical))

                    reasons = set(reasons_base)
                    if link.code_conflict:
                        reasons.add("access_code_conflict")

                    if kind == "main" and slug is not None:
                        main_pairs.add((slug, link.canonical))

                    if link_key not in agg.links:
                        _ensure_media(agg, identity, info, m_type)
                        _ensure_group(agg, group_key, edition)
                        agg.links[link_key] = {
                            "provider": link.provider,
                            "canonical": link.canonical,
                            "url": link.url,
                            "label": link.label,
                            "access_code": link.access_code,
                            "group_key": group_key,
                            "title_raw": title_cell.strip() if title_cell else None,
                            "remark": remark_cell.strip() if remark_cell else None,
                            "created_at_list": [created_at] if created_at is not None else [],
                            "deleted_at_list": [deleted_at],
                        }
                        if kind != "main":
                            agg.dedup[f"{kind}_added"] += 1
                        # The report's reason counts are scoped to distinct
                        # links (one count per newly-created link), matching
                        # how totals/providers are also counted over the
                        # deduped result rather than raw source rows: a
                        # duplicate occurrence of an already-known link
                        # (e.g. a supplement row copying a main-table row)
                        # must not inflate year_missing/needs_review again.
                        for reason in reasons:
                            agg.needs_review[reason] += 1
                        if "access_code_conflict" in reasons:
                            agg.conflicts["access_code_conflict"] += 1
                        if "year_missing" in reasons:
                            agg.year_missing += 1
                        if reasons:
                            _mark_review(agg, group_key, identity, reasons)
                    else:
                        existing = agg.links[link_key]
                        # Minor: first wins (existing["access_code"] is left
                        # untouched) but a repeated canonical link that
                        # carries a different access code across rows is
                        # now counted.
                        if link.access_code != existing["access_code"]:
                            agg.conflicts["access_code_divergence"] += 1
                        if created_at is not None:
                            existing["created_at_list"].append(created_at)
                        existing["deleted_at_list"].append(deleted_at)

                        if kind == "main":
                            agg.dedup["within_main_duplicate_links"] += 1
                        else:
                            agg.dedup[f"{kind}_matched"] += 1

                        if existing["group_key"] != group_key:
                            agg.conflicts["link_shared_across_groups"] += 1
                            agg.needs_review["link_shared_across_groups"] += 1
                            _ensure_media(agg, identity, info, m_type)
                            _ensure_group(agg, group_key, edition)
                            _mark_review(agg, group_key, identity, {"link_shared_across_groups"})
                            old_identity = existing["group_key"][0]
                            _mark_review(agg, existing["group_key"], old_identity, {"link_shared_across_groups"})
                            if reasons:
                                _mark_review(agg, group_key, identity, reasons)
                        elif reasons:
                            _mark_review(agg, group_key, identity, reasons)

                    agg.provenance.append({
                        "link_key": link_key, "source_file": filename, "sheet": sheet,
                        "row_number": row_number, "record_id": record_id, "slug": slug,
                        "owner_uid": owner_uid, "owner_tag": owner_tag,
                    })
            except ImportRowError:
                stat["rows_error"] += 1
                continue

        agg.sources[kind] = stat

    # §4.8 rule 1: cross-check the official 磁力ED2K sheet against the main
    # table. Non-影视 rows (the sheet includes 追更 rows too) are not
    # eligible for import in the first place, so they are excluded here.
    cc_sheet = CROSS_CHECK[1]
    cross_rows = rows_by_source.get("cross_check", [])
    missing = 0
    for row in cross_rows:
        if _cell_str(row.get("类型")) != "影视":
            continue
        slug_raw = row.get("Slug/追更包Slug")
        slug = str(slug_raw).strip() if slug_raw not in (None, "") else None
        if slug is None:
            # Same guard as main_pairs' population above: a row with no
            # Slug was never added to main_pairs either, so comparing
            # (None, canonical) here would always miss even when the main
            # row it corresponds to also has no Slug -- an empty Slug on
            # both sides must never produce a false "missing" pair.
            continue
        # I8: same no-leak conversion the main aggregation loop uses -- but,
        # unlike a main-loop failure (I7), NOT contained: a broken
        # cross-check row means §4.8 rule 1 can no longer be verified at
        # all, so this must still escape aggregate() and be caught by
        # run()'s own try/except, producing the same minimal failure report
        # and exit code 1 as any other pre-aggregate() failure.
        cc_link_cell = _cell_str(row.get("链接")) or ""
        cc_code_cell = _cell_str(row.get("访问码"))
        try:
            link = parse_link(cc_link_cell, cc_code_cell)
        except Exception:
            raise ImportRowError(cc_sheet, row.get("_row_number", 0), "链接") from None
        if (slug, link.canonical) not in main_pairs:
            missing += 1
    agg.dedup["cross_check_ed2k_sheet_subset"] = (missing == 0)
    agg.dedup["cross_check_missing_pairs"] = missing

    for link_key, entry in agg.links.items():
        entry["created_at_source"] = min(entry["created_at_list"]) if entry["created_at_list"] else None
        deleted_list = entry["deleted_at_list"]
        all_deleted = bool(deleted_list) and all(d is not None for d in deleted_list)
        entry["deleted_at_source"] = max(deleted_list) if all_deleted else None
        agg.providers[link_key[0]] += 1

        # media.created_at/resource_group.created_at must also be
        # deterministic (derived from the source data, never wall-clock),
        # for the same idempotency reason as resource_link.imported_at
        # above: reuse the earliest created_at_source seen among each
        # media's/group's own links.
        if entry["created_at_source"] is not None:
            media_identity_, _fingerprint = entry["group_key"]
            for holder in (agg.groups.get(entry["group_key"]), agg.media.get(media_identity_)):
                if holder is None:
                    continue
                if holder.get("created_at") is None or entry["created_at_source"] < holder["created_at"]:
                    holder["created_at"] = entry["created_at_source"]

    for media in agg.media.values():
        agg.media_type[media["media_type"]] += 1

    agg.totals = {
        "media": len(agg.media),
        "resource_groups": len(agg.groups),
        "links": len(agg.links),
        "provenance_rows": len(agg.provenance),
    }
    return agg


# ---------------------------------------------------------------------------
# write_bundle
# ---------------------------------------------------------------------------


def _default_pinyin(text: str) -> tuple:
    from pypinyin import lazy_pinyin

    syllables = lazy_pinyin(text or "")
    full = "".join(syllables)
    initials = "".join(s[0] for s in syllables if s)
    return full, initials


def _default_charmap_convert(ch: str) -> str:
    from zhconv import convert

    return convert(ch, "zh-hans")


# ---------------------------------------------------------------------------
# T15 (z1-hints §design item 1 + master work order v1.4 §14.3): Codex's
# offline IMDb candidate hints and (optional) official IMDb ratings, both
# shipped as gzipped JSONL artifacts alongside the library bundle. Neither
# ever carries a resource URL, share code or credential -- only titles,
# years, media types and IMDb ids, per docs/metadata-enrichment.md.
# ---------------------------------------------------------------------------

_HINT_CANDIDATE_FIELDS = (
    "imdb_id", "imdb_type", "start_year", "end_year", "primary_title", "original_title", "score", "reasons",
)
_MAX_HINT_CANDIDATES = 3


def _open_jsonl(path: Path):
    """Open a ``.jsonl`` or ``.jsonl.gz`` file for text reading, by
    extension."""
    if str(path).endswith(".gz"):
        return gzip.open(path, "rt", encoding="utf-8")
    return open(path, "rt", encoding="utf-8")


def _sanitize_hint_candidate(raw: dict) -> dict:
    """Keep only the design item 1 allowlisted fields -- never an evidence
    blob or a URL, even if a future producer accidentally includes one."""
    return {key: raw[key] for key in _HINT_CANDIDATE_FIELDS if key in raw}


@dataclass
class HintsResult:
    """Outcome of ``load_hints()``: counts for the importer report (§design
    item 1: "Report counts by decision") plus the rows actually eligible to
    write into the bundle's ``tmdb_hints`` table."""

    rows_total: int = 0
    by_decision: Counter = field(default_factory=Counter)
    unmatched_identity: int = 0
    imdb_ids: set = field(default_factory=set)
    matched: list = field(default_factory=list)  # (media_identity, source, decision, candidates)


def load_hints(hints_path: Path, media_identities: set, *, source: str = "imdb_offline") -> HintsResult:
    """Read Codex's offline IMDb candidate hints (one JSON object per line;
    see docs/metadata-enrichment.md §2 for the row
    shape) and match each row to this import's own freshly-aggregated
    ``media_identity`` set -- NOT the row's own ``media_id`` (an
    unmatched-media staging id from the read-only export that has no
    meaning here). ``media_identity`` is computed with the exact §design
    item 1 formula: ``library_normalize.media_identity(TitleInfo(title_zh=
    row['title_zh'], aliases=(), original_hint=None, year=row.get('year')),
    row['media_type'], None)``. A row whose identity isn't in
    ``media_identities`` is counted (``unmatched_identity``) and skipped --
    it is never written to ``tmdb_hints``. At most the top 3 candidates
    (already ranked by the producer) are kept, each trimmed to the
    allowlisted fields only."""
    result = HintsResult()
    with _open_jsonl(hints_path) as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            result.rows_total += 1
            decision = row.get("decision") or ""
            result.by_decision[decision] += 1

            info = TitleInfo(title_zh=row.get("title_zh") or "", aliases=(), original_hint=None, year=row.get("year"))
            identity = media_identity(info, row.get("media_type") or "unknown", None)
            if identity not in media_identities:
                result.unmatched_identity += 1
                continue

            candidates = [_sanitize_hint_candidate(c) for c in (row.get("candidates") or [])[:_MAX_HINT_CANDIDATES]]
            for candidate in candidates:
                if candidate.get("imdb_id"):
                    result.imdb_ids.add(candidate["imdb_id"])
            result.matched.append((identity, source, decision, candidates))
    return result


def write_hints(store: "LibraryStore", hints: HintsResult, *, generated_at: int) -> None:
    """Write ``hints.matched`` rows into the bundle's ``tmdb_hints`` table
    (keyed by ``media_identity``, the schema's own PRIMARY KEY -- a repeat
    import simply replaces the row)."""
    if not hints.matched:
        return
    conn = store.connect()
    try:
        for identity, source, decision, candidates in hints.matched:
            conn.execute(
                "INSERT INTO tmdb_hints (media_identity, source, decision, candidate_count, candidates_json, generated_at) "
                "VALUES (?, ?, ?, ?, ?, ?) "
                "ON CONFLICT(media_identity) DO UPDATE SET source=excluded.source, decision=excluded.decision, "
                "candidate_count=excluded.candidate_count, candidates_json=excluded.candidates_json, "
                "generated_at=excluded.generated_at",
                (identity, source, decision, len(candidates), json.dumps(candidates, ensure_ascii=False), generated_at),
            )
        conn.commit()
    finally:
        conn.close()


@dataclass
class ImdbRatingsResult:
    rows_loaded: int = 0
    file_found: bool = False


def load_imdb_ratings(ratings_path: Path, relevant_ids: set, *, as_of: str) -> tuple[list, ImdbRatingsResult]:
    """Read the (optional) official-IMDb-ratings artifact (one
    ``{"imdb_id": "tt...", "rating": 8.7, "votes": 31198}`` object per
    line, see docs/metadata-enrichment.md) and
    keep only rows whose ``imdb_id`` is in ``relevant_ids`` (the ids
    referenced by this import's own loaded hints and confirmations --
    never the whole dataset, §14.3). Skips gracefully (empty result,
    ``file_found=False``)
    when the artifact doesn't exist yet -- it is expected to ship later."""
    result = ImdbRatingsResult()
    if not relevant_ids or not Path(ratings_path).exists():
        return [], result
    result.file_found = True
    rows: list = []
    with _open_jsonl(ratings_path) as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            imdb_id = row.get("imdb_id")
            if imdb_id in relevant_ids:
                rows.append((imdb_id, row.get("rating"), row.get("votes"), as_of))
    result.rows_loaded = len(rows)
    return rows, result


def write_imdb_ratings(store: "LibraryStore", rows: list) -> None:
    if not rows:
        return
    conn = store.connect()
    try:
        conn.executemany(
            "INSERT INTO imdb_ratings (imdb_id, rating, votes, as_of) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(imdb_id) DO UPDATE SET rating=excluded.rating, votes=excluded.votes, as_of=excluded.as_of",
            rows,
        )
        conn.commit()
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# w5-confirm-import: Codex's manual TMDB-identity confirmations (see
# docs/metadata-enrichment.md §1-§6). Unlike
# load_hints() above (which matches by media_identity, recomputed from the
# row's own title/year/type), a confirmation row is keyed by the media_id a
# read-only export of a *previous* build of this same bundle assigned --
# so it must be validated against, and mapped through, the media_id this
# run's own write_bundle() will assign to the same identity.
# ---------------------------------------------------------------------------

_CONFIRMATION_EVIDENCE_FIELDS = ("year_exact", "media_type_exact", "wikidata_tmdb_property", "sources")
_CONFIRMATION_SKIP_REASONS = ("missing", "title_mismatch", "year_mismatch", "type_mismatch", "bad_row", "duplicate")


def _media_lookup_from_aggregate(agg: "Aggregate") -> dict:
    """Map the ``media_id`` ``write_bundle()`` will assign to each media row
    -- 1-based, in ``sorted(agg.media)`` order, exactly the order
    ``write_bundle()`` calls ``store.upsert_media()`` in, so a fresh rebuild
    from the same (unchanged) source workbooks reproduces the same ids a
    previous build (the one Codex's confirmations file was validated
    against) already assigned -- back to the attributes
    ``load_confirmations()`` needs to validate a confirmation row against."""
    lookup: dict = {}
    for media_id, identity in enumerate(sorted(agg.media), start=1):
        m = agg.media[identity]
        lookup[media_id] = {
            "media_identity": identity,
            "title_zh": m["title_zh"],
            "year": m["year"],
            "media_type": m["media_type"],
        }
    return lookup


def _sanitize_confirmation_evidence(raw) -> dict:
    """Keep only the §contract allowlisted evidence fields -- never an
    arbitrary key a future producer might accidentally include."""
    if not isinstance(raw, dict):
        return {}
    return {key: raw[key] for key in _CONFIRMATION_EVIDENCE_FIELDS if key in raw}


@dataclass
class ConfirmationsResult:
    """Outcome of ``load_confirmations()``: counts for the importer report
    plus the rows actually eligible to write into ``tmdb_hints`` as
    ``confirmed`` hints."""

    rows_total: int = 0
    skipped: Counter = field(default_factory=Counter)
    accepted: list = field(default_factory=list)  # (media_identity, candidate)
    shared_tmdb_pairs: int = 0

    @property
    def written(self) -> int:
        return len(self.accepted)


def load_confirmations(path: Path, media_lookup: dict) -> ConfirmationsResult:
    """Read Codex's offline manual-confirmation rows (one JSON object per
    line; see docs/metadata-enrichment.md §4) and
    validate each against ``media_lookup`` (``_media_lookup_from_aggregate``
    of this run's own freshly-aggregated media): ``media_id`` present AND
    ``title_zh`` identical AND ``(year or 0)`` identical AND ``media_type``
    identical -> accepted; otherwise skipped, counted by reason. A row that
    can't even be parsed, or is missing/mistypes a required field, counts as
    ``bad_row`` rather than raising. A ``media_id`` that has already
    produced an accepted row is a ``duplicate`` -- only the first accepted
    row per ``media_id`` is kept, so ``written`` (``len(accepted)``) always
    equals the number of distinct media_ids written."""
    result = ConfirmationsResult()
    accepted_media_ids: set = set()
    with _open_jsonl(path) as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            result.rows_total += 1
            try:
                row = json.loads(line)
                media_id = row["media_id"]
                title_zh = row["title_zh"]
                year = row.get("year")
                media_type = row["media_type"]
                tmdb_id = row["tmdb_id"]
                if media_type not in ("movie", "tv"):
                    raise ValueError("bad media_type")
                if not isinstance(tmdb_id, int) or isinstance(tmdb_id, bool):
                    raise ValueError("bad tmdb_id")
            except (KeyError, ValueError, TypeError, json.JSONDecodeError):
                result.skipped["bad_row"] += 1
                continue

            if media_id in accepted_media_ids:
                result.skipped["duplicate"] += 1
                continue

            entry = media_lookup.get(media_id)
            if entry is None:
                result.skipped["missing"] += 1
                continue
            if entry["title_zh"] != title_zh:
                result.skipped["title_mismatch"] += 1
                continue
            if (entry["year"] or 0) != (year or 0):
                result.skipped["year_mismatch"] += 1
                continue
            if entry["media_type"] != media_type:
                result.skipped["type_mismatch"] += 1
                continue

            candidate = {
                "imdb_id": row.get("imdb_id"),
                "tmdb_id": tmdb_id,
                "tmdb_type": media_type,
                "confidence": row.get("confidence"),
                "original_decision": row.get("original_decision"),
                "evidence": _sanitize_confirmation_evidence(row.get("evidence")),
            }
            result.accepted.append((entry["media_identity"], candidate))
            accepted_media_ids.add(media_id)

    pair_counts = Counter((candidate["tmdb_type"], candidate["tmdb_id"]) for _identity, candidate in result.accepted)
    result.shared_tmdb_pairs = sum(1 for count in pair_counts.values() if count > 1)
    return result


def write_confirmations(store: "LibraryStore", result: "ConfirmationsResult", *, generated_at: int) -> None:
    """Write ``result.accepted`` rows into the bundle's ``tmdb_hints`` table
    as ``confirmed`` hints (``source='codex_manual_confirmation'``,
    ``candidate_count=1``), replacing any existing hint for that identity --
    called after ``write_hints()`` so confirmed rows always win. Each
    candidate records the ``decision`` of the hint row it replaces (``None``
    if there wasn't one yet) as ``prior_hint_decision``, read from the table
    just before it gets overwritten."""
    if not result.accepted:
        return
    conn = store.connect()
    try:
        for identity, candidate in result.accepted:
            existing = conn.execute(
                "SELECT decision FROM tmdb_hints WHERE media_identity=?", (identity,)
            ).fetchone()
            prior_hint_decision = existing[0] if existing else None
            full_candidate = {**candidate, "prior_hint_decision": prior_hint_decision}
            candidates_json = json.dumps([full_candidate], ensure_ascii=False)
            conn.execute(
                "INSERT INTO tmdb_hints (media_identity, source, decision, candidate_count, candidates_json, generated_at) "
                "VALUES (?, 'codex_manual_confirmation', 'confirmed', 1, ?, ?) "
                "ON CONFLICT(media_identity) DO UPDATE SET source=excluded.source, decision=excluded.decision, "
                "candidate_count=excluded.candidate_count, candidates_json=excluded.candidates_json, "
                "generated_at=excluded.generated_at",
                (identity, candidates_json, generated_at),
            )
        conn.commit()
    finally:
        conn.close()


def write_bundle(
    agg: Aggregate,
    output_dir: Path,
    *,
    source_hashes: dict,
    pinyin=None,
    charmap_convert=None,
    hints: "HintsResult | None" = None,
    imdb_ratings_rows: list | None = None,
    confirmations: "ConfirmationsResult | None" = None,
) -> Path:
    """Write a fresh ``library-bundle.sqlite`` from ``agg`` and return its
    path.  Always rebuilds from zero (a stale ``.tmp`` is discarded first),
    written under ``umask 0077`` then atomically ``os.replace``'d into
    place with explicit ``0600`` permissions.

    ``hints``/``imdb_ratings_rows`` (T15), when given, are written into the
    same bundle file via ``write_hints``/``write_imdb_ratings`` -- part of
    the same atomic tmp-then-replace write as everything else here.
    ``confirmations`` (w5-confirm-import), when given, is written via
    ``write_confirmations`` after ``write_hints`` so a confirmed identity
    always wins over a plain offline hint for the same media.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    # SQLite's temp files (VACUUM/index builds) go next to the output, not
    # to /tmp -- on a tiny tmpfs the import dies with "database or disk is
    # full" (the same failure install_bundle guards against).
    _use_temp_dir(output_dir)
    tmp_path = output_dir / "library-bundle.sqlite.tmp"
    final_path = output_dir / "library-bundle.sqlite"

    for suffix in ("", "-wal", "-shm", "-journal"):
        stray = Path(str(tmp_path) + suffix)
        if stray.exists():
            stray.unlink()

    old_umask = os.umask(0o077)
    try:
        store = LibraryStore(tmp_path)
        store.create_schema()

        media_id_map: dict = {}
        for identity in sorted(agg.media):
            m = agg.media[identity]
            rec = MediaRecord(
                media_identity=identity,
                media_type=m["media_type"],
                title_zh=m["title_zh"],
                search_key=m["search_key"],
                title_original=m["title_original"],
                title_alt_json=json.dumps(sorted(m["aliases"]), ensure_ascii=False),
                year=m["year"],
                match_status="needs_review" if m["needs_review"] else "unmatched",
                created_at=m.get("created_at") or 0,
                updated_at=m.get("created_at") or 0,
            )
            media_id_map[identity] = store.upsert_media(rec)

        group_id_map: dict = {}
        for group_key in sorted(agg.groups):
            g = agg.groups[group_key]
            rec = GroupRecord(
                media_id=media_id_map[group_key[0]],
                edition_fingerprint=group_key[1],
                display_title=g["display_title"],
                season_from=g["season_from"], season_to=g["season_to"],
                episode_from=g["episode_from"], episode_to=g["episode_to"],
                complete_season=int(bool(g["complete_season"])),
                quality=g["quality"], source_type=g["source_type"], hdr=g["hdr"],
                video_codec=g["video_codec"],
                audio_summary="+".join(g["audio"]) or None,
                subtitle_summary="+".join(g["subtitle"]) or None,
                tags_json=json.dumps(list(g["tags"]), ensure_ascii=False),
                needs_review=int(bool(g["needs_review"])),
                review_reason=",".join(sorted(g["review_reasons"])) or None,
                created_at=g.get("created_at") or 0,
                updated_at=g.get("created_at") or 0,
            )
            group_id_map[group_key] = store.upsert_group(rec)

        link_id_map: dict = {}
        for link_key in sorted(agg.links):
            l = agg.links[link_key]  # noqa: E741
            rec = LinkRecord(
                public_id=public_id(l["provider"], l["canonical"]),
                group_id=group_id_map[l["group_key"]],
                provider=l["provider"],
                canonical_url_hash=link_key[1],
                url_label=l["label"],
                url_plain=l["url"],
                access_code_plain=l["access_code"],
                has_access_code=int(bool(l["access_code"])),
                title_raw=l["title_raw"],
                remark=l["remark"],
                created_at_source=l["created_at_source"],
                deleted_at_source=l["deleted_at_source"],
                # Deterministic, not wall-clock: two builds from the same
                # input must produce byte-for-byte identical rows except
                # schema_meta.built_at (see write_bundle's idempotency
                # contract), so this can't default to time.time().
                imported_at=l["created_at_source"] if l["created_at_source"] is not None else 0,
            )
            link_id, _created = store.upsert_link(rec)
            link_id_map[link_key] = link_id

        for entry in sorted(agg.provenance, key=lambda p: (p["source_file"], p["sheet"], p["row_number"])):
            store.add_provenance(ProvenanceRecord(
                link_id=link_id_map[entry["link_key"]],
                source_file=entry["source_file"], sheet=entry["sheet"],
                row_number=entry["row_number"], record_id=entry["record_id"],
                slug=entry["slug"], owner_uid=entry["owner_uid"], owner_tag=entry["owner_tag"],
            ))

        store.recount()

        pinyin_fn = pinyin if pinyin is not None else _default_pinyin
        convert_fn = charmap_convert if charmap_convert is not None else _default_charmap_convert

        texts: list = []
        for m in agg.media.values():
            texts.append(m["title_zh"])
            texts.extend(m["aliases"])
            if m["title_original"]:
                texts.append(m["title_original"])

        from library_search import build_charmap, build_index

        charmap_map = build_charmap(texts, convert_fn)
        build_index(store, pinyin=pinyin_fn, charmap=charmap_map)

        generated_at = int(time.time())
        if hints is not None:
            write_hints(store, hints, generated_at=generated_at)
        if imdb_ratings_rows:
            write_imdb_ratings(store, imdb_ratings_rows)
        if confirmations is not None:
            write_confirmations(store, confirmations, generated_at=generated_at)

        store.meta_set("normalize_version", NORMALIZE_VERSION)
        store.meta_set("source_hashes", json.dumps(source_hashes, sort_keys=True))
        store.meta_set("built_at", str(int(time.time())))
        store.meta_set("encrypted", "0")

        conn = store.connect()
        try:
            conn.execute("PRAGMA wal_checkpoint(TRUNCATE)")
            conn.execute("PRAGMA journal_mode=DELETE")
            conn.commit()
        finally:
            conn.close()

        for suffix in ("-wal", "-shm", "-journal"):
            stray = Path(str(tmp_path) + suffix)
            if stray.exists():
                stray.unlink()

        os.replace(tmp_path, final_path)
        os.chmod(final_path, 0o600)
        return final_path
    finally:
        os.umask(old_umask)


# ---------------------------------------------------------------------------
# write_report
# ---------------------------------------------------------------------------


def _hints_report_dict(hints_result: "HintsResult | None") -> dict | None:
    if hints_result is None:
        return None
    return {
        "rows_total": hints_result.rows_total,
        "by_decision": dict(hints_result.by_decision),
        "matched": len(hints_result.matched),
        "unmatched_identity": hints_result.unmatched_identity,
        "distinct_imdb_ids": len(hints_result.imdb_ids),
    }


def _imdb_ratings_report_dict(ratings_result: "ImdbRatingsResult | None") -> dict | None:
    if ratings_result is None:
        return None
    return {"file_found": ratings_result.file_found, "rows_loaded": ratings_result.rows_loaded}


def _confirmations_report_dict(result: "ConfirmationsResult | None") -> dict | None:
    if result is None:
        return None
    return {
        "rows": result.rows_total,
        "written": result.written,
        "skipped": {reason: result.skipped.get(reason, 0) for reason in _CONFIRMATION_SKIP_REASONS},
        "shared_tmdb_pairs": result.shared_tmdb_pairs,
    }


def _build_report_dict(
    agg: Aggregate, sources_meta: list, mode: str, *,
    hints_result: "HintsResult | None" = None, imdb_ratings_report: "ImdbRatingsResult | None" = None,
    confirmations_result: "ConfirmationsResult | None" = None,
) -> dict:
    sources = []
    for meta in sources_meta:
        stat = agg.sources.get(meta["kind"], {})
        sources.append({
            "file": meta["file"], "sheet": meta["sheet"],
            "sha256": meta.get("sha256"), "header_row": meta.get("header_row"),
            "rows_total": stat.get("rows_total", 0),
            "rows_video": stat.get("rows_video", 0),
            "rows_skipped_non_video": stat.get("rows_skipped_non_video", 0),
            "rows_type_missing": stat.get("rows_type_missing", 0),
            "rows_invalid_link": stat.get("rows_invalid_link", 0),
            "rows_missing_link": stat.get("rows_missing_link", 0),
            "rows_deleted": stat.get("rows_deleted", 0),
            "rows_bad_timestamp": stat.get("rows_bad_timestamp", 0),
            "rows_shifted": stat.get("rows_shifted", 0),
            "rows_missing_record_id": stat.get("rows_missing_record_id", 0),
            "rows_multi_link": stat.get("rows_multi_link", 0),
            "links_from_multi_link_cells": stat.get("links_from_multi_link_cells", 0),
            "rows_error": stat.get("rows_error", 0),
        })
    dedup = {k: v for k, v in agg.dedup.items() if k != "cross_check_missing_pairs"}
    report = {
        "version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "normalize_version": NORMALIZE_VERSION,
        "mode": mode,
        "status": "ok",
        "sources": sources,
        "totals": dict(agg.totals),
        "providers": dict(agg.providers),
        "media_type": dict(agg.media_type),
        "year_missing": agg.year_missing,
        "dedup": dedup,
        "conflicts": dict(agg.conflicts),
        "needs_review": dict(agg.needs_review),
        "elapsed_seconds": agg.elapsed_seconds,
    }
    hints_dict = _hints_report_dict(hints_result)
    if hints_dict is not None:
        report["hints"] = hints_dict
    ratings_dict = _imdb_ratings_report_dict(imdb_ratings_report)
    if ratings_dict is not None:
        report["imdb_ratings"] = ratings_dict
    confirmations_dict = _confirmations_report_dict(confirmations_result)
    if confirmations_dict is not None:
        report["confirmations"] = confirmations_dict
    return report


def write_report(
    agg: Aggregate, sources_meta: list, path: Path, mode: str, *,
    hints_result: "HintsResult | None" = None, imdb_ratings_report: "ImdbRatingsResult | None" = None,
    confirmations_result: "ConfirmationsResult | None" = None,
) -> None:
    """Write the desensitised JSON report (§5) -- counts and hashes only,
    never a link, access code or title."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    report = _build_report_dict(
        agg, sources_meta, mode, hints_result=hints_result, imdb_ratings_report=imdb_ratings_report,
        confirmations_result=confirmations_result,
    )
    old_umask = os.umask(0o077)
    try:
        path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    finally:
        os.umask(old_umask)
    os.chmod(path, 0o600)


def _write_failure_report(path: Path | None, mode: str, error_cls: type) -> None:
    """A minimal report for a failure that happens before ``aggregate()``
    ever runs (missing source file, ``--require-hashes`` mismatch, or an
    ``ImportRowError``) -- just enough for a caller to see that the run
    failed and what kind of failure it was.  Never any text that could
    carry a cell value (sheet/column names, titles, links...); a no-op
    when ``path`` is None (no ``--report`` was requested)."""
    if path is None:
        return
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "mode": mode,
        "status": "failed",
        "error": error_cls.__name__,
    }
    old_umask = os.umask(0o077)
    try:
        path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    finally:
        os.umask(old_umask)
    os.chmod(path, 0o600)


# ---------------------------------------------------------------------------
# main / CLI
# ---------------------------------------------------------------------------


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


_README_HASH_RE = re.compile(r"^\|\s*`([^`]+)`\s*\|\s*`([0-9a-fA-F]+)`\s*\|", re.MULTILINE)


def _read_reference_hashes(readme_path: Path) -> dict:
    if not readme_path.is_file():
        return {}
    text = readme_path.read_text(encoding="utf-8")
    return {name: value.lower() for name, value in _README_HASH_RE.findall(text)}


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="import_media_library.py",
        description="Offline Excel importer for the personal media-resource library.",
    )
    parser.add_argument("--source-dir", dest="source_dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--write", action="store_true", default=False)
    parser.add_argument("--report", type=Path, default=_DEFAULT_REPORT_PATH)
    parser.add_argument("--require-hashes", dest="require_hashes", action="store_true", default=False)
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument(
        "--hints", type=Path, default=None,
        help="optional Codex offline IMDb candidate hints (.jsonl or .jsonl.gz), keyed by media_identity (T15)",
    )
    parser.add_argument(
        "--imdb-ratings", dest="imdb_ratings", type=Path, default=None,
        help="optional official-IMDb-ratings artifact (.jsonl or .jsonl.gz); skipped gracefully if missing (T15)",
    )
    parser.add_argument(
        "--confirmations", type=Path, default=None,
        help="optional Codex manual TMDB-identity confirmations (.jsonl or .jsonl.gz), keyed by media_id (w5-confirm-import)",
    )
    return parser


def run(args: argparse.Namespace) -> tuple:
    """Run one full import pass. Returns ``(report_dict_or_None, exit_code)``."""
    start = time.time()
    source_dir = Path(args.source_dir)
    mode = "write" if args.write else "dry-run"

    source_hashes: dict = {}
    hash_mismatches: list = []
    reference_hashes = _read_reference_hashes(source_dir / "README.md")
    all_filenames = [name for name, _sheet, _kind in SOURCES]
    for filename in all_filenames:
        file_path = source_dir / filename
        if not file_path.is_file():
            print(f"error: source file not found: {filename}", file=sys.stderr)
            _write_failure_report(args.report, mode, SourceFileMissingError)
            return None, 1
        digest = _sha256_file(file_path)
        source_hashes[filename] = digest
        expected = reference_hashes.get(filename)
        if expected is not None and expected != digest:
            hash_mismatches.append(filename)

    if hash_mismatches:
        print(
            "warning: source hash mismatch for: " + ", ".join(sorted(hash_mismatches)),
            file=sys.stderr,
        )
        if args.require_hashes:
            _write_failure_report(args.report, mode, SourceHashMismatchError)
            return None, 1

    rows_by_source: dict = {}
    sources_meta: list = []
    try:
        for filename, sheet, kind in SOURCES:
            header_row, rows = read_sheet(source_dir / filename, sheet)
            if args.limit is not None:
                rows = rows[: args.limit]
            rows_by_source[kind] = rows
            sources_meta.append({
                "file": filename, "sheet": sheet, "kind": kind,
                "sha256": source_hashes[filename], "header_row": header_row,
            })

        cc_file, cc_sheet = CROSS_CHECK
        _cc_header_row, cc_rows = read_sheet(source_dir / cc_file, cc_sheet)
        if args.limit is not None:
            cc_rows = cc_rows[: args.limit]
        rows_by_source["cross_check"] = cc_rows

        # aggregate() can also raise ImportRowError (e.g. a cell parse_title/
        # parse_edition/parse_link/_parse_int can't handle) -- it
        # must be caught here too, not just read_sheet()'s errors above, so
        # such a failure still produces the minimal failure report and exit
        # code 1 instead of escaping main() uncaught.
        agg = aggregate(rows_by_source)
    except ImportRowError as exc:
        print(f"error: {exc}", file=sys.stderr)
        _write_failure_report(args.report, mode, type(exc))
        return None, 1

    # T15: hints/ratings are matched against this run's own freshly-
    # aggregated media_identity set regardless of --write, so a dry-run
    # still reports decision counts and identity mismatches (the report
    # deliverable this task's brief asks for) without ever needing a bundle.
    hints_result: HintsResult | None = None
    imdb_ratings_report: ImdbRatingsResult | None = None
    imdb_ratings_rows: list = []
    if args.hints is not None:
        hints_result = load_hints(args.hints, set(agg.media))

    # w5-confirm-import: confirmations are matched against this run's own
    # freshly-aggregated media (via _media_lookup_from_aggregate), same as
    # hints/ratings above -- a dry-run still reports counts without a bundle.
    # Loaded here, before the --imdb-ratings step below, so a confirmed
    # identity's own imdb_id counts as "relevant" too -- not just the ids
    # --hints happened to also propose as a candidate.
    confirmations_result: ConfirmationsResult | None = None
    if args.confirmations is not None:
        media_lookup = _media_lookup_from_aggregate(agg)
        confirmations_result = load_confirmations(args.confirmations, media_lookup)

    if args.hints is not None and args.imdb_ratings is not None:
        as_of = datetime.now(timezone.utc).date().isoformat()
        relevant_ids = set(hints_result.imdb_ids)
        if confirmations_result is not None:
            relevant_ids.update(
                candidate["imdb_id"]
                for _identity, candidate in confirmations_result.accepted
                if candidate.get("imdb_id")
            )
        imdb_ratings_rows, imdb_ratings_report = load_imdb_ratings(args.imdb_ratings, relevant_ids, as_of=as_of)

    # A cross-check failure (§4.8 rule 1) is known immediately after
    # aggregate() returns -- skip write_bundle entirely rather than writing
    # a bundle and then having to discard it, so a failed --write run never
    # leaves one behind.
    if args.write and agg.dedup.get("cross_check_ed2k_sheet_subset", True):
        write_bundle(
            agg, args.output, source_hashes=source_hashes,
            hints=hints_result, imdb_ratings_rows=imdb_ratings_rows,
            confirmations=confirmations_result,
        )

    agg.elapsed_seconds = time.time() - start

    if args.report is not None:
        write_report(
            agg, sources_meta, args.report, mode, hints_result=hints_result, imdb_ratings_report=imdb_ratings_report,
            confirmations_result=confirmations_result,
        )

    report = _build_report_dict(
        agg, sources_meta, mode, hints_result=hints_result, imdb_ratings_report=imdb_ratings_report,
        confirmations_result=confirmations_result,
    )

    if mode == "write" and not agg.dedup.get("cross_check_ed2k_sheet_subset", True):
        return report, 1
    return report, 0


def main(argv: list | None = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)
    report, code = run(args)
    if report is not None:
        print(json.dumps(report, ensure_ascii=False))
    return code


if __name__ == "__main__":
    raise SystemExit(main())
