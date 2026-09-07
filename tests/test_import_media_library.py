"""Tests for scripts/import_media_library.py (T2.5): offline Excel importer.

Two layers:

* Rule-level unit tests call ``read_sheet``/``aggregate`` directly against
  small hand-built ``rows_by_source`` dicts (no real workbook needed) so each
  rule in the brief (§4.8) gets one focused, fast test.
* An end-to-end test runs the full pipeline (``read_sheet`` -> ``aggregate``
  -> ``write_bundle`` -> ``write_report``, plus ``main()``) against the
  synthetic ``library_sources`` fixture and checks the counts against
  ``EXPECTED`` from ``fixtures/library/make_workbooks.py`` (computed there
  from the same row literals, never hand-typed).

All links/hostnames below are the fixtures' fictional values (``swfake``/
``fake`` share codes, 4-character access codes) per the project's test-data
rule; nothing here is a real share link or access code.
"""

from __future__ import annotations

import json
import sqlite3
import stat
import sys
from datetime import datetime
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import library_store as ls  # noqa: E402
import import_media_library as importer  # noqa: E402
from fixtures.library.make_workbooks import (  # noqa: E402
    EXPECTED,
    SUPPLEMENT_ED2K_ROWS,
    SUPPLEMENT_VIDEO_ROWS,
    VIDEO_ROWS,
)

# ---------------------------------------------------------------------------
# Helpers to build synthetic rows_by_source dicts for the rule-level tests.
# read_sheet() attaches an internal "_row_number" key to every row dict; the
# rule-level tests build that shape directly instead of going through a real
# workbook.
# ---------------------------------------------------------------------------


def _row(
    *,
    type_="影视",
    record_id=1,
    slug="slug-1",
    uid=500001,
    badge="无能",
    title="示例剧 (2024)",
    media_title="示例剧",
    link="https://115cdn.com/s/swfakeaa11",
    code="",
    remark="S01 4K WEB-DL DV 内封简繁",
    created=datetime(2024, 1, 1),
    deleted=None,
    row_number=5,
):
    return {
        "类型": type_,
        "记录ID": record_id,
        "Slug/追更包Slug": slug,
        "用户UID": uid,
        "官组标记": badge,
        "标题": title,
        "媒体标题": media_title,
        "链接": link,
        "访问码": code,
        "备注": remark,
        "创建时间": created,
        "删除时间": deleted,
        "_row_number": row_number,
    }


def _rows_by_source(main=None, supplement_video=None, supplement_ed2k=None, cross_check=None):
    return {
        "main": main or [],
        "supplement_video": supplement_video or [],
        "supplement_ed2k": supplement_ed2k or [],
        "cross_check": cross_check or [],
    }


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------


def test_header_and_cross_check_constants_match_brief():
    assert importer.HEADER == [
        "类型", "记录ID", "Slug/追更包Slug", "用户UID", "官组标记", "标题",
        "媒体标题", "链接", "访问码", "备注", "创建时间", "删除时间",
    ]
    assert importer.SOURCES == [
        ("HDHive official-group-backup.xlsx", "影视分享", "main"),
        ("影巢_影视分享_整理.xlsx", "影视分享", "supplement_video"),
        ("影巢_ED2K_整理.xlsx", "磁力ED2K", "supplement_ed2k"),
    ]
    assert importer.CROSS_CHECK == ("HDHive official-group-backup.xlsx", "磁力ED2K")


# ---------------------------------------------------------------------------
# read_sheet
# ---------------------------------------------------------------------------


def test_read_sheet_finds_header_row_and_skips_blank_record_id(library_sources):
    header_row, rows = importer.read_sheet(
        library_sources / "HDHive official-group-backup.xlsx", "影视分享"
    )
    assert header_row == 4
    assert len(rows) == len(VIDEO_ROWS)
    assert rows[0]["类型"] == "影视"
    assert rows[0]["_row_number"] == 5


def test_read_sheet_row_dict_has_all_header_keys(library_sources):
    _header_row, rows = importer.read_sheet(
        library_sources / "HDHive official-group-backup.xlsx", "影视分享"
    )
    for name in importer.HEADER:
        assert name in rows[0]


def test_read_sheet_missing_header_column_raises_import_row_error(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "影视分享"
    # header row missing "访问码"
    bad_header = [h for h in importer.HEADER if h != "访问码"]
    for col, name in enumerate(bad_header, start=1):
        ws.cell(row=1, column=col, value=name)
    ws.cell(row=2, column=1, value=1)
    path = tmp_path / "bad.xlsx"
    wb.save(path)

    with pytest.raises(importer.ImportRowError) as excinfo:
        importer.read_sheet(path, "影视分享")
    assert excinfo.value.sheet == "影视分享"
    assert excinfo.value.column == "访问码"


def test_read_sheet_missing_record_id_header_raises_import_row_error(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "影视分享"
    ws.cell(row=1, column=1, value="not a header row")
    path = tmp_path / "no-header.xlsx"
    wb.save(path)

    with pytest.raises(importer.ImportRowError):
        importer.read_sheet(path, "影视分享")


def test_read_sheet_header_tolerates_bom_and_trailing_spaces(tmp_path):
    """Minor: a header cell with a trailing space or a leading BOM ("﻿",
    which Excel/CSV-round-tripped workbooks sometimes leave on the very
    first header cell) must still be recognised as its HEADER name -- data
    rows are unaffected."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "影视分享"
    noisy_header = [f"﻿{importer.HEADER[0]}"] + [f"{name} " for name in importer.HEADER[1:]]
    for col, name in enumerate(noisy_header, start=1):
        ws.cell(row=1, column=col, value=name)
    ws.cell(row=2, column=1, value="影视")
    ws.cell(row=2, column=2, value=1)
    path = tmp_path / "bom-header.xlsx"
    wb.save(path)

    header_row, rows = importer.read_sheet(path, "影视分享")
    assert header_row == 1
    assert len(rows) == 1
    assert rows[0]["类型"] == "影视"
    assert rows[0]["记录ID"] == 1


# ---------------------------------------------------------------------------
# ImportRowError never carries cell values
# ---------------------------------------------------------------------------


def test_import_row_error_class_itself_never_carries_the_cell_value():
    """ImportRowError's constructor only ever takes sheet/row_number/column
    -- never the cell value -- so this holds structurally regardless of
    where the exception is raised or caught."""
    bad_cell_marker = "not-a-real-record-id-zz99-fake-marker"
    exc = importer.ImportRowError("影视分享", 42, "记录ID")
    assert exc.sheet == "影视分享"
    assert exc.row_number == 42
    assert exc.column == "记录ID"
    assert bad_cell_marker not in str(exc)
    assert bad_cell_marker not in repr(exc)


def _agg_dump(agg) -> str:
    """Stringify everything an Aggregate could ever have accumulated (media/
    groups/links/provenance/counters), for a no-leak scan -- write_report()
    only surfaces a subset of this, but a leak into any of it would
    eventually reach write_bundle()'s callers or a future report field.
    (``repr()``, not ``json.dumps()``: several dict keys here are tuples,
    e.g. (provider, canonical_hash), which json can't key on.)"""
    return repr(agg.__dict__)


def test_bad_record_id_row_is_contained_not_fatal_and_never_leaks_the_cell_value():
    """I7: an unparseable 记录ID no longer aborts the whole run -- it is
    contained as sources[].rows_error, the row is skipped (not stored), and
    the run continues; this checks that containment never leaks the bad
    cell value into any state aggregate() keeps."""
    bad_cell_marker = "not-a-real-record-id-zz99-fake-marker"
    good_row = _row(record_id=1, slug="s1", link="https://pan.quark.cn/s/fakegoodrid1", row_number=5)
    bad_row = _row(record_id=bad_cell_marker, slug="s-bad", row_number=42)
    agg = importer.aggregate(_rows_by_source(main=[good_row, bad_row]))
    assert agg.sources["main"]["rows_error"] == 1
    assert agg.totals["links"] == 1
    assert bad_cell_marker not in _agg_dump(agg)


def test_aggregate_parse_link_failure_is_contained_and_never_leaks_cell_value(monkeypatch):
    """I7: a parse_link failure on one row is contained the same way -- one
    good row plus one row whose link value triggers the (simulated) failure
    still imports the good row's link, and the bad row's value never
    reaches any state aggregate() keeps. Hotfix regression this also
    covers: on the real workbooks, parse_link's own ValueError (raised for
    some malformed ED2K cells before the parse_link hotfix) was letting the
    raw cell value escape through aggregate() uncaught."""
    secret = "SECRET-CELL-VALUE-fake-marker"
    real_parse_link = importer.parse_link

    def boom(raw_url, access_code_cell):
        if raw_url == secret:
            raise ValueError(secret)
        return real_parse_link(raw_url, access_code_cell)

    monkeypatch.setattr(importer, "parse_link", boom)

    good_row = _row(record_id=1, slug="s1", link="https://pan.quark.cn/s/fakegoodlink1", row_number=6)
    bad_row = _row(record_id=2, slug="s2", link=secret, row_number=7)
    agg = importer.aggregate(_rows_by_source(main=[good_row, bad_row]))

    assert agg.sources["main"]["rows_error"] == 1
    assert agg.totals["links"] == 1
    assert secret not in _agg_dump(agg)


def test_cross_check_parse_link_failure_is_not_contained_and_propagates(monkeypatch):
    """I8: unlike a main-loop failure (I7), a parse_link failure in the
    cross-check loop is NOT contained -- it must still surface as
    ImportRowError (never a raw/leaky exception) and propagate out of
    aggregate(), so run() can still write the failure report and exit 1."""
    trigger = "TRIGGER-CROSS-CHECK-BOOM-fake-marker"
    real_parse_link = importer.parse_link

    def boom(raw_url, access_code_cell):
        if raw_url == trigger:
            raise ValueError("boom")
        return real_parse_link(raw_url, access_code_cell)

    monkeypatch.setattr(importer, "parse_link", boom)

    main_row = _row(record_id=1, slug="s1", row_number=5)
    cross_row = _row(record_id=1, slug="s1", link=trigger, row_number=9)

    with pytest.raises(importer.ImportRowError) as excinfo:
        importer.aggregate(_rows_by_source(main=[main_row], cross_check=[cross_row]))
    exc = excinfo.value
    assert exc.sheet == importer.CROSS_CHECK[1]
    assert exc.row_number == 9
    assert exc.column == "链接"
    assert trigger not in str(exc)
    assert trigger not in repr(exc)


# ---------------------------------------------------------------------------
# parse_source_timestamp -- real-data hotfix: the 影巢 supplement workbooks
# hold 创建时间/删除时间 as strings (e.g. "2025-05-08 15:35:13"), not
# datetime objects like the main workbook.
# ---------------------------------------------------------------------------


def test_parse_source_timestamp_accepts_datetime():
    dt = datetime(2025, 5, 8, 15, 35, 13)
    assert importer.parse_source_timestamp(dt) == int(dt.timestamp())


def test_parse_source_timestamp_accepts_space_separated_string():
    dt = datetime(2025, 5, 8, 15, 35, 13)
    assert importer.parse_source_timestamp("2025-05-08 15:35:13") == int(dt.timestamp())


def test_parse_source_timestamp_accepts_t_separated_string():
    dt = datetime(2025, 5, 8, 15, 35, 13)
    assert importer.parse_source_timestamp("2025-05-08T15:35:13") == int(dt.timestamp())


def test_parse_source_timestamp_accepts_slash_separated_string():
    dt = datetime(2025, 5, 8, 15, 35, 13)
    assert importer.parse_source_timestamp("2025/05/08 15:35:13") == int(dt.timestamp())


def test_parse_source_timestamp_accepts_date_only_string():
    dt = datetime(2025, 5, 8)
    assert importer.parse_source_timestamp("2025-05-08") == int(dt.timestamp())


def test_parse_source_timestamp_is_case_and_whitespace_tolerant():
    dt = datetime(2025, 5, 8, 15, 35, 13)
    assert importer.parse_source_timestamp("  2025-05-08t15:35:13  ") == int(dt.timestamp())
    assert importer.parse_source_timestamp("  2025-05-08  ") == int(datetime(2025, 5, 8).timestamp())


def test_parse_source_timestamp_none_and_empty_string_are_none():
    assert importer.parse_source_timestamp(None) is None
    assert importer.parse_source_timestamp("") is None
    assert importer.parse_source_timestamp("   ") is None


def test_parse_source_timestamp_garbage_string_is_none():
    assert importer.parse_source_timestamp("not-a-real-date-zz99-fake-marker") is None


# ---------------------------------------------------------------------------
# aggregate() wiring: a garbage 创建时间/删除时间 string never raises -- it
# is counted in sources[].rows_bad_timestamp and flagged needs_review
# "timestamp_unparsed" instead (real-data hotfix).
# ---------------------------------------------------------------------------


def test_bad_created_timestamp_string_does_not_raise_and_is_counted():
    row = _row(record_id=1, slug="s1", created="not-a-real-date-zz99-fake-marker")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_bad_timestamp"] == 1
    assert agg.needs_review["timestamp_unparsed"] == 1
    (entry,) = agg.links.values()
    assert entry["created_at_source"] is None


def test_bad_deleted_timestamp_string_is_also_counted():
    row = _row(record_id=1, slug="s1", deleted="not-a-real-date-zz99-fake-marker")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_bad_timestamp"] == 1
    assert agg.needs_review["timestamp_unparsed"] == 1
    # Minor: rows_deleted counts a 删除时间 that actually parsed, not the
    # cell's raw truthiness -- this string is truthy but unparseable.
    assert agg.sources["main"]["rows_deleted"] == 0


def test_rows_deleted_counts_only_a_parsed_deletion_timestamp():
    good_row = _row(record_id=1, slug="s1", deleted=datetime(2024, 6, 1), row_number=5)
    bad_row = _row(
        record_id=2, slug="s2", link="https://pan.quark.cn/s/fakebaddel1",
        deleted="not-a-real-date-zz99-fake-marker", row_number=6,
    )
    agg = importer.aggregate(_rows_by_source(main=[good_row, bad_row]))
    assert agg.sources["main"]["rows_deleted"] == 1
    assert agg.sources["main"]["rows_bad_timestamp"] == 1


def test_good_timestamp_string_does_not_trigger_bad_timestamp_counter():
    row = _row(record_id=1, slug="s1", created="2024-01-01 00:00:00")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_bad_timestamp"] == 0
    assert agg.needs_review["timestamp_unparsed"] == 0


def test_string_created_timestamp_matches_equivalent_datetime_created_timestamp():
    """The two 影巢 workbooks hold 创建时间 as a string; the resulting
    created_at_source must be identical to what the same instant would
    produce from the main workbook's datetime cells."""
    row = _row(record_id=1, slug="s1", created="2025-05-08 15:35:13")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    (entry,) = agg.links.values()
    assert entry["created_at_source"] == int(datetime(2025, 5, 8, 15, 35, 13).timestamp())


# ---------------------------------------------------------------------------
# Shifted-row repair (real-data hotfix, aggregate()): 12 rows in the main
# workbook are shifted one column right starting at 链接 -- 链接 is empty,
# 访问码 holds the URL, 备注 holds the (optional) access code, 创建时间
# holds the remark text, 删除时间 holds the real creation datetime, and the
# real deletion time was never captured by the shift (always lost -> None).
# These rows are built by hand (not via _row(), which only produces normal
# rows) to get the shifted column layout exactly.
# ---------------------------------------------------------------------------


def test_shifted_rows_are_repaired_and_normal_invalid_link_row_still_counts():
    shifted_with_code = {
        "类型": "影视", "记录ID": 1, "Slug/追更包Slug": "slug-shift-1",
        "用户UID": 500001, "官组标记": "无能",
        "标题": "错位剧甲 (2024)", "媒体标题": "错位剧甲",
        "链接": None, "访问码": "https://115cdn.com/s/swfakeshift1",
        "备注": "ab12", "创建时间": "S01 4K WEB-DL",
        "删除时间": datetime(2024, 1, 1), "_row_number": 5,
    }
    shifted_without_code = {
        "类型": "影视", "记录ID": 2, "Slug/追更包Slug": "slug-shift-2",
        "用户UID": 500001, "官组标记": "无能",
        "标题": "错位剧乙 (2024)", "媒体标题": "错位剧乙",
        "链接": None, "访问码": "https://pan.quark.cn/s/fakeshift2",
        "备注": None, "创建时间": "S01 4K WEB-DL",
        "删除时间": datetime(2024, 2, 1), "_row_number": 6,
    }
    normal_row = _row(record_id=3, slug="slug-normal", row_number=7)
    # An empty 链接 whose 访问码 does NOT look like a link (the existing
    # test-data convention: empty string) must still not be mistaken for a
    # shifted row -- it is now a genuine blank-link row (I9), counted in
    # rows_missing_link and skipped, not an "invalid link" (rows_invalid_link
    # is for a non-blank cell that fails to parse as any known scheme; see
    # blank_link_row below and the dedicated invalid-link test).
    blank_link_row = _row(record_id=4, slug="slug-blank", link="", row_number=8)
    invalid_link_row = _row(record_id=5, slug="slug-invalid", link="请私信要资源", row_number=9)

    agg = importer.aggregate(_rows_by_source(
        main=[shifted_with_code, shifted_without_code, normal_row, blank_link_row, invalid_link_row]
    ))

    assert agg.sources["main"]["rows_shifted"] == 2
    assert agg.needs_review["row_shifted"] == 2
    assert agg.sources["main"]["rows_missing_link"] == 1
    assert agg.sources["main"]["rows_invalid_link"] == 1

    entries_by_url = {e["url"]: e for e in agg.links.values() if e["url"]}

    entry_a = entries_by_url["https://115cdn.com/s/swfakeshift1"]
    assert entry_a["provider"] == "115"
    assert entry_a["access_code"] == "ab12"
    assert entry_a["remark"] == "S01 4K WEB-DL"
    assert entry_a["created_at_source"] == int(datetime(2024, 1, 1).timestamp())
    assert entry_a["deleted_at_source"] is None

    entry_b = entries_by_url["https://pan.quark.cn/s/fakeshift2"]
    assert entry_b["access_code"] is None
    assert entry_b["remark"] == "S01 4K WEB-DL"
    assert entry_b["created_at_source"] == int(datetime(2024, 2, 1).timestamp())
    assert entry_b["deleted_at_source"] is None


def test_shifted_row_access_code_longer_than_eight_chars_is_dropped():
    shifted_row = {
        "类型": "影视", "记录ID": 1, "Slug/追更包Slug": "slug-shift-3",
        "用户UID": 500001, "官组标记": "无能",
        "标题": "错位剧丙 (2024)", "媒体标题": "错位剧丙",
        "链接": None, "访问码": "https://pan.quark.cn/s/fakeshift3",
        "备注": "toolongcode123", "创建时间": "S01 4K WEB-DL",
        "删除时间": datetime(2024, 3, 1), "_row_number": 9,
    }
    agg = importer.aggregate(_rows_by_source(main=[shifted_row]))
    (entry,) = agg.links.values()
    assert entry["access_code"] is None


# ---------------------------------------------------------------------------
# I9/I10: blank -- or, once stripped, whitespace-only -- 链接 rows are
# counted in rows_missing_link and skipped entirely, never stored (never
# routed through parse_link at all, so they can never collapse into a
# shared "unknown:<sha256 of \"\">" link or conflict with each other).
# ---------------------------------------------------------------------------


def test_blank_link_row_is_skipped_and_counted_missing_link():
    row = _row(record_id=1, slug="s1", link="")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_missing_link"] == 1
    assert agg.totals["links"] == 0
    assert agg.provenance == []


def test_two_blank_link_rows_do_not_collapse_or_conflict():
    rows = [
        _row(record_id=1, slug="s1", link="", row_number=5),
        _row(record_id=2, slug="s2", link="", row_number=6),
    ]
    agg = importer.aggregate(_rows_by_source(main=rows))
    assert agg.sources["main"]["rows_missing_link"] == 2
    assert agg.totals["links"] == 0
    assert agg.conflicts["access_code_conflict"] == 0
    assert agg.conflicts["link_shared_across_groups"] == 0


def test_whitespace_only_link_is_treated_as_blank_and_counted_missing_link():
    """I10: a whitespace-only 链接 cell must be stripped before the
    blank-link truthiness test, not treated as a "present" (garbage) link."""
    row = _row(record_id=1, slug="s1", link="   ")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_missing_link"] == 1
    assert agg.totals["links"] == 0


def test_whitespace_only_link_does_not_defeat_shift_detection():
    """I10: a whitespace-only (not merely absent) 链接 cell must still let
    _is_row_shifted recognise a shifted row -- stripping happens before the
    truthiness check there too."""
    shifted_row = {
        "类型": "影视", "记录ID": 1, "Slug/追更包Slug": "slug-shift-ws",
        "用户UID": 500001, "官组标记": "无能",
        "标题": "错位剧丁 (2024)", "媒体标题": "错位剧丁",
        "链接": "   ", "访问码": "https://pan.quark.cn/s/fakeshiftws1",
        "备注": "ab12", "创建时间": "S01 4K WEB-DL",
        "删除时间": datetime(2024, 1, 1), "_row_number": 10,
    }
    agg = importer.aggregate(_rows_by_source(main=[shifted_row]))
    assert agg.sources["main"]["rows_shifted"] == 1
    (entry,) = agg.links.values()
    assert entry["url"] == "https://pan.quark.cn/s/fakeshiftws1"
    assert entry["access_code"] == "ab12"


# ---------------------------------------------------------------------------
# §4.8 rule 2: 类型 != 影视 skipped; I12: 类型 compared after str(...).strip(),
# with a None/blank cell counted separately (rows_type_missing) from a real
# non-影视 type (rows_skipped_non_video).
# ---------------------------------------------------------------------------


def test_non_video_rows_are_skipped_and_counted():
    rows = [
        _row(record_id=1, slug="s1", type_="音乐", link="https://pan.quark.cn/s/fakem1"),
        _row(record_id=2, slug="s2", type_="影视", link="https://pan.quark.cn/s/fakev1"),
    ]
    agg = importer.aggregate(_rows_by_source(main=rows))
    assert agg.sources["main"]["rows_skipped_non_video"] == 1
    assert agg.sources["main"]["rows_video"] == 1
    assert agg.totals["links"] == 1


def test_type_cell_with_trailing_whitespace_still_counts_as_video():
    row = _row(record_id=1, slug="s1", type_=" 影视 ")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_video"] == 1
    assert agg.sources["main"]["rows_skipped_non_video"] == 0
    assert agg.totals["links"] == 1


def test_blank_type_cell_is_counted_separately_from_non_video():
    rows = [
        _row(record_id=1, slug="s1", type_=None, link="https://pan.quark.cn/s/faketype1"),
        _row(record_id=2, slug="s2", type_="音乐", link="https://pan.quark.cn/s/faketype2", row_number=6),
    ]
    agg = importer.aggregate(_rows_by_source(main=rows))
    assert agg.sources["main"]["rows_type_missing"] == 1
    assert agg.sources["main"]["rows_skipped_non_video"] == 1
    assert agg.totals["links"] == 0


# ---------------------------------------------------------------------------
# §4.8 rule 3: invalid link -> provider unknown, still stored
# ---------------------------------------------------------------------------


def test_invalid_link_becomes_unknown_provider_and_is_still_stored():
    row = _row(record_id=1, slug="s1", link="请私信要资源，站内联系")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_invalid_link"] == 1
    assert agg.totals["links"] == 1
    (link_key,) = agg.links.keys()
    assert link_key[0] == "unknown"


# ---------------------------------------------------------------------------
# I11: a 链接 cell can pack several links back to back (real data: 2-142
# ED2K links in one cell, no reliable separator) -- split on a lookahead
# for each recognised scheme so every one becomes its own link, sharing the
# row's title/edition and provenance (sheet/row_number).
# ---------------------------------------------------------------------------


def test_multi_link_cell_with_three_links_becomes_three_links():
    ed2k_1 = "ed2k://|file|Fake.Multi.A|111111|00000000000000000000000000000001|/"
    ed2k_2 = "ed2k://|file|Fake.Multi.B|222222|00000000000000000000000000000002|/"
    https_1 = "https://pan.quark.cn/s/fakemulti03"
    packed = ed2k_1 + ed2k_2 + https_1
    row = _row(record_id=1, slug="s1", link=packed, row_number=11)

    agg = importer.aggregate(_rows_by_source(main=[row]))

    assert agg.sources["main"]["rows_multi_link"] == 1
    assert agg.sources["main"]["links_from_multi_link_cells"] == 3
    assert agg.totals["links"] == 3
    assert len(agg.provenance) == 3
    for entry in agg.provenance:
        assert entry["sheet"] == "影视分享"
        assert entry["row_number"] == 11
    urls = {e["url"] for e in agg.links.values()}
    assert urls == {ed2k_1, ed2k_2, https_1}


def test_single_link_cell_is_unaffected_by_multi_link_split():
    row = _row(record_id=1, slug="s1", link="https://pan.quark.cn/s/fakesingle01")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_multi_link"] == 0
    assert agg.sources["main"]["links_from_multi_link_cells"] == 0
    assert agg.totals["links"] == 1


# ---------------------------------------------------------------------------
# §4.8 rule 4: dedup key, provenance-only duplicate, created/deleted rules
# ---------------------------------------------------------------------------


def test_duplicate_link_only_appends_provenance_not_new_media_or_group():
    same_link = "https://www.guangyapan.com/s/1234efgh"
    rows = [
        _row(record_id=101, slug="s101", link=same_link, row_number=5),
        _row(record_id=102, slug="s102", link=same_link, row_number=6),
    ]
    agg = importer.aggregate(_rows_by_source(main=rows))
    assert agg.totals["links"] == 1
    assert agg.totals["media"] == 1
    assert agg.totals["resource_groups"] == 1
    assert agg.totals["provenance_rows"] == 2
    assert agg.dedup["within_main_duplicate_links"] == 1


def test_record_id_and_slug_are_provenance_only_not_dedup_key():
    """Two rows with different 记录ID/Slug but the same link/title/edition
    dedup to one link -- record_id/slug never participate in the dedup key."""
    same_link = "https://pan.quark.cn/s/fakedup1"
    rows = [
        _row(record_id=9001, slug="slug-a", link=same_link),
        _row(record_id=9002, slug="slug-b", link=same_link),
    ]
    agg = importer.aggregate(_rows_by_source(main=rows))
    assert agg.totals["links"] == 1
    slugs = {p["slug"] for p in agg.provenance}
    record_ids = {p["record_id"] for p in agg.provenance}
    assert slugs == {"slug-a", "slug-b"}
    assert record_ids == {9001, 9002}


def test_blank_record_id_row_is_counted_and_not_skipped():
    """Minor: a blank/None 记录ID -- which read_sheet() itself would never
    hand to aggregate() (it skips such rows at the source), but a row-level
    unit test can still build directly -- is counted in
    rows_missing_record_id without being dropped: the row's link is still
    imported, just with record_id=None in its provenance."""
    row = _row(record_id="", slug="s1", link="https://pan.quark.cn/s/fakenorid1")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.sources["main"]["rows_missing_record_id"] == 1
    assert agg.totals["links"] == 1
    (entry,) = agg.provenance
    assert entry["record_id"] is None


def test_created_at_source_takes_the_earliest_occurrence():
    same_link = "https://pan.quark.cn/s/fakeearliest1"
    rows = [
        _row(record_id=1, slug="s1", link=same_link, created=datetime(2024, 3, 1)),
        _row(record_id=2, slug="s2", link=same_link, created=datetime(2024, 1, 1)),
    ]
    agg = importer.aggregate(_rows_by_source(main=rows))
    (entry,) = agg.links.values()
    assert entry["created_at_source"] == int(datetime(2024, 1, 1).timestamp())


def test_deleted_at_source_requires_every_provenance_row_deleted():
    same_link = "https://www.guangyapan.com/s/1234efgh"
    rows = [
        _row(record_id=1, slug="s1", link=same_link, deleted=None),
        _row(record_id=2, slug="s2", link=same_link, deleted=datetime(2024, 6, 1)),
    ]
    agg = importer.aggregate(_rows_by_source(main=rows))
    (entry,) = agg.links.values()
    assert entry["deleted_at_source"] is None, "not all provenance rows are deleted"


def test_deleted_at_source_set_when_all_provenance_rows_deleted():
    same_link = "https://www.guangyapan.com/s/1234efgh"
    rows = [
        _row(record_id=1, slug="s1", link=same_link, deleted=datetime(2024, 5, 1)),
        _row(record_id=2, slug="s2", link=same_link, deleted=datetime(2024, 6, 1)),
    ]
    agg = importer.aggregate(_rows_by_source(main=rows))
    (entry,) = agg.links.values()
    assert entry["deleted_at_source"] == int(datetime(2024, 6, 1).timestamp())


def test_link_shared_across_groups_conflict_flags_both_groups():
    same_link = "https://pan.quark.cn/s/fakeconflict1"
    row_a = _row(
        record_id=1, slug="s1", link=same_link,
        title="冲突剧 (2024)", media_title="冲突剧", remark="S01 4K WEB-DL",
    )
    row_b = _row(
        record_id=2, slug="s2", link=same_link,
        title="冲突剧 (2024)", media_title="冲突剧", remark="S02 4K WEB-DL",
    )
    agg = importer.aggregate(_rows_by_source(main=[row_a, row_b]))

    assert agg.conflicts["link_shared_across_groups"] == 1
    assert agg.needs_review["link_shared_across_groups"] == 1
    # the link itself did not move: it still belongs to row_a's group
    (entry,) = agg.links.values()
    assert agg.totals["resource_groups"] == 2, "group B is created just to be flagged"
    review_flags = [g["needs_review"] for g in agg.groups.values()]
    assert review_flags == [True, True]


def test_access_code_divergence_counter_first_wins():
    """Minor: a repeated canonical link carrying a different access code
    across rows keeps first-wins semantics (the stored access_code is the
    first row's), but the divergence is now counted."""
    same_link = "https://pan.quark.cn/s/fakediv1"
    row_a = _row(record_id=1, slug="s1", link=same_link, code="ab12", row_number=5)
    row_b = _row(record_id=2, slug="s2", link=same_link, code="cd34", row_number=6)
    agg = importer.aggregate(_rows_by_source(main=[row_a, row_b]))

    assert agg.conflicts["access_code_divergence"] == 1
    (entry,) = agg.links.values()
    assert entry["access_code"] == "ab12", "first wins"


def test_no_access_code_divergence_when_codes_match():
    same_link = "https://pan.quark.cn/s/fakediv2"
    row_a = _row(record_id=1, slug="s1", link=same_link, code="ab12", row_number=5)
    row_b = _row(record_id=2, slug="s2", link=same_link, code="ab12", row_number=6)
    agg = importer.aggregate(_rows_by_source(main=[row_a, row_b]))
    assert agg.conflicts["access_code_divergence"] == 0


# ---------------------------------------------------------------------------
# §4.8 rule 6: needs_review reasons
# ---------------------------------------------------------------------------


def test_year_missing_triggers_needs_review_and_is_counted():
    row = _row(record_id=1, slug="s1", title="无年份剧", media_title="无年份剧")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.needs_review["year_missing"] == 1
    assert agg.year_missing == 1
    (media,) = agg.media.values()
    assert media["needs_review"] is True


def test_title_alias_conflict_triggers_needs_review():
    row = _row(
        record_id=1, slug="s1",
        title="旧城计划 (2019)", media_title="旧城重制计划",
    )
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.needs_review["title_alias_conflict"] == 1


def test_title_alias_substring_is_not_a_conflict():
    row = _row(
        record_id=1, slug="s1",
        title="旧日回声 / 往昔低语 (2022)", media_title="旧日回声",
    )
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.needs_review["title_alias_conflict"] == 0


def test_edition_unparsed_triggers_needs_review():
    row = _row(record_id=1, slug="s1", remark="随便写的备注")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.needs_review["edition_unparsed"] == 1


def test_empty_remark_is_not_edition_unparsed():
    row = _row(record_id=1, slug="s1", remark="")
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.needs_review["edition_unparsed"] == 0


def test_access_code_conflict_triggers_needs_review_and_conflicts():
    row = _row(
        record_id=1, slug="s1",
        link="https://115.com/s/swfakerowk1?password=ab12", code="zz99",
    )
    agg = importer.aggregate(_rows_by_source(main=[row]))
    assert agg.needs_review["access_code_conflict"] == 1
    assert agg.conflicts["access_code_conflict"] == 1


# ---------------------------------------------------------------------------
# §4.8 rule 1: main-table rows sorted by (创建时间 asc, None last), then
# 记录ID asc, before aggregation -- so the earliest-created row always wins
# a link's first-occurrence fields, regardless of the order rows_by_source
# happens to hand them in.
# ---------------------------------------------------------------------------


def test_main_rows_out_of_order_earliest_created_row_wins_first_occurrence():
    same_link = "https://pan.quark.cn/s/fakeorder1"
    row_early = _row(
        record_id=1, slug="s1", link=same_link,
        title="示例剧 (2024)", media_title="示例剧",
        remark="S01 4K WEB-DL", created=datetime(2024, 1, 1), row_number=5,
    )
    row_late = _row(
        record_id=2, slug="s2", link=same_link,
        title="示例剧一 (2024)", media_title="示例剧",
        # "补" (resend) does not affect edition_fingerprint (§4.7), so this
        # row shares row_early's group_key while still being observably
        # different from it -- a real way to tell which row "won".
        remark="S01 4K WEB-DL 补", created=datetime(2024, 6, 1), row_number=6,
    )
    # Later-created row listed FIRST in the input.
    agg = importer.aggregate(_rows_by_source(main=[row_late, row_early]))

    (entry,) = agg.links.values()
    assert entry["title_raw"] == "示例剧 (2024)"
    assert entry["remark"] == "S01 4K WEB-DL"
    (group,) = agg.groups.values()
    assert group["tags"] == (), "group edition should come from the earliest row (no resend tag)"


def test_main_rows_aggregate_output_is_order_independent_after_sort():
    same_link = "https://pan.quark.cn/s/fakeorder2"
    row_early = _row(
        record_id=1, slug="s1", link=same_link,
        title="示例剧 (2024)", media_title="示例剧",
        remark="S01 4K WEB-DL", created=datetime(2024, 1, 1), row_number=5,
    )
    row_late = _row(
        record_id=2, slug="s2", link=same_link,
        title="示例剧一 (2024)", media_title="示例剧",
        remark="S01 4K WEB-DL 补", created=datetime(2024, 6, 1), row_number=6,
    )
    agg_forward = importer.aggregate(_rows_by_source(main=[row_early, row_late]))
    agg_reversed = importer.aggregate(_rows_by_source(main=[row_late, row_early]))
    assert agg_forward == agg_reversed


# ---------------------------------------------------------------------------
# Cross-check (§4.8 rule 1) -- one passing and one FAILING case.
# ---------------------------------------------------------------------------


def test_cross_check_passes_when_ed2k_sheet_is_subset_of_main():
    link = "ed2k://|file|Fake.Show.S01E01.2160p|123456|0123456789ABCDEF0123456789ABCDEF01|/"
    main_row = _row(record_id=1, slug="slug-ed2k-1", type_="影视", link=link)
    cross_row = _row(record_id=1, slug="slug-ed2k-1", type_="影视", link=link)
    agg = importer.aggregate(_rows_by_source(main=[main_row], cross_check=[cross_row]))
    assert agg.dedup["cross_check_ed2k_sheet_subset"] is True


def test_cross_check_fails_when_ed2k_sheet_has_a_row_not_in_main():
    link_in_main = "ed2k://|file|Fake.Show.S01E01.2160p|123456|0123456789ABCDEF0123456789ABCDEF01|/"
    link_not_in_main = "ed2k://|file|Fake.Show.NeverImported|999999|0123456789ABCDEF0123456789ABCDEFFF|/"
    main_row = _row(record_id=1, slug="slug-ed2k-1", type_="影视", link=link_in_main)
    cross_row_ok = _row(record_id=1, slug="slug-ed2k-1", type_="影视", link=link_in_main)
    cross_row_orphan = _row(record_id=2, slug="slug-ed2k-2", type_="影视", link=link_not_in_main)
    agg = importer.aggregate(
        _rows_by_source(main=[main_row], cross_check=[cross_row_ok, cross_row_orphan])
    )
    assert agg.dedup["cross_check_ed2k_sheet_subset"] is False


def test_cross_check_ignores_non_video_rows_in_the_ed2k_sheet():
    """The official 磁力ED2K sheet includes 追更-type rows (per the fixture
    generator's comment); those are not eligible for import in the first
    place, so they must not fail the subset check."""
    link_in_main = "ed2k://|file|Fake.Show.S01E01.2160p|123456|0123456789ABCDEF0123456789ABCDEF01|/"
    zhuigeng_link = "ed2k://|file|Fake.Zhuigeng.Pack|999999|0123456789ABCDEF0123456789ABCDEFAA|/"
    main_row = _row(record_id=1, slug="slug-ed2k-1", type_="影视", link=link_in_main)
    cross_row_ok = _row(record_id=1, slug="slug-ed2k-1", type_="影视", link=link_in_main)
    cross_row_zhuigeng = _row(record_id=1, slug="slug-zhuigeng", type_="追更", link=zhuigeng_link)
    agg = importer.aggregate(
        _rows_by_source(main=[main_row], cross_check=[cross_row_ok, cross_row_zhuigeng])
    )
    assert agg.dedup["cross_check_ed2k_sheet_subset"] is True


def test_cross_check_skips_rows_with_no_slug_on_either_side():
    """An empty Slug on both the main row and its corresponding official
    磁力ED2K row must never register as a "missing" pair: main_pairs never
    contains a None-slug entry (Slug is provenance-only, §4.8 rule 5), so
    the cross-check side must apply the identical skip -- not compare
    (None, canonical) against a bucket that was never populated for it."""
    link = "ed2k://|file|Fake.Show.NoSlug|123456|0123456789ABCDEF0123456789ABCDEF09|/"
    main_row = _row(record_id=1, slug=None, type_="影视", link=link)
    cross_row = _row(record_id=1, slug=None, type_="影视", link=link)
    agg = importer.aggregate(_rows_by_source(main=[main_row], cross_check=[cross_row]))
    assert agg.dedup["cross_check_ed2k_sheet_subset"] is True


# ---------------------------------------------------------------------------
# Supplement matched/added dedup counters
# ---------------------------------------------------------------------------


def test_supplement_video_matched_vs_added():
    main_link = "https://115cdn.com/s/swfakeaa11"
    new_link = "https://115.com/s/swfakenew1"
    main_row = _row(record_id=1, slug="s1", link=main_link)
    supp_matched = _row(record_id=1, slug="s1", link=main_link)
    supp_added = _row(record_id=9001, slug="s9001", link=new_link, title="新剧 (2021)", media_title="新剧")
    agg = importer.aggregate(
        _rows_by_source(main=[main_row], supplement_video=[supp_matched, supp_added])
    )
    assert agg.dedup["supplement_video_matched"] == 1
    assert agg.dedup["supplement_video_added"] == 1
    assert agg.totals["links"] == 2


# ---------------------------------------------------------------------------
# edition_display_title (§3 resource_group.display_title -- review fix 1:
# it must be the edition summary, never the media title)
# ---------------------------------------------------------------------------


def test_edition_display_title_unspecified_edition():
    from library_normalize import parse_edition

    edition = parse_edition("", "")
    assert importer.edition_display_title(edition) == "未标注版本"


def test_edition_display_title_season_quality_source_hdr_subtitle():
    from library_normalize import parse_edition

    edition = parse_edition("S01 4K WEB-DL DV 内封简繁", "")
    assert importer.edition_display_title(edition) == "S01 · 4K · WEB-DL · DV · 内封简繁"


def test_edition_display_title_episode_range_and_alt_subtitle_phrase():
    from library_normalize import parse_edition

    # "内嵌简中" folds to the same internal subtitle code as "内封简中"
    # (§4.3); the display form must be the canonical phrase, not whichever
    # source text happened to be typed.
    edition = parse_edition("S01E01 - E06 1080P 内嵌简中 仅秒传", "")
    assert importer.edition_display_title(edition) == "S01E01-E06 · 1080p · 内封简中"


# ---------------------------------------------------------------------------
# write_bundle / write_report
# ---------------------------------------------------------------------------


def _small_agg():
    rows = [
        _row(
            record_id=1, slug="s1", link="https://115cdn.com/s/swfakeaa11",
            code="ab12", title="示例剧 (2024)", media_title="示例剧",
        ),
        _row(
            record_id=2, slug="s2", link="https://pan.quark.cn/s/fakeaa12",
            title="示例剧二 (2023)", media_title="示例剧二", row_number=6,
        ),
    ]
    return importer.aggregate(_rows_by_source(main=rows))


def test_write_bundle_creates_sqlite_with_0600_permissions(tmp_path):
    agg = _small_agg()
    out_dir = tmp_path / "out"
    result = importer.write_bundle(agg, out_dir, source_hashes={"a.xlsx": "deadbeef"})
    assert result == out_dir / "library-bundle.sqlite"
    assert result.is_file()
    mode = stat.S_IMODE(result.stat().st_mode)
    assert mode == 0o600
    # no stray -wal/-shm/.tmp files left behind
    leftovers = sorted(p.name for p in out_dir.iterdir())
    assert leftovers == ["library-bundle.sqlite"], leftovers


def test_write_bundle_content_matches_aggregate(tmp_path):
    agg = _small_agg()
    path = importer.write_bundle(agg, tmp_path / "out", source_hashes={})
    conn = sqlite3.connect(str(path))
    try:
        media_count = conn.execute("SELECT COUNT(*) FROM media").fetchone()[0]
        link_count = conn.execute("SELECT COUNT(*) FROM resource_link").fetchone()[0]
        prov_count = conn.execute("SELECT COUNT(*) FROM link_provenance").fetchone()[0]
        # plaintext link/access-code values must never be masked away in the
        # store itself (the store is where the real data belongs; only
        # reports/exceptions must avoid them) but they must round-trip.
        url = conn.execute("SELECT url_plain FROM resource_link WHERE provider='115'").fetchone()[0]
    finally:
        conn.close()
    assert media_count == agg.totals["media"]
    assert link_count == agg.totals["links"]
    assert prov_count == agg.totals["provenance_rows"]
    assert url == "https://115cdn.com/s/swfakeaa11"


def test_write_bundle_is_idempotent_across_two_builds(tmp_path):
    agg = _small_agg()
    out_dir = tmp_path / "out"
    path_a = importer.write_bundle(agg, out_dir, source_hashes={"a.xlsx": "deadbeef"})
    dump_a = _dump_tables(path_a)
    path_b = importer.write_bundle(agg, out_dir, source_hashes={"a.xlsx": "deadbeef"})
    dump_b = _dump_tables(path_b)
    assert dump_a == dump_b


def _dump_tables(path: Path) -> dict:
    tables = [
        "media", "resource_group", "resource_link", "link_provenance",
        "search_doc", "search_term", "search_vocab", "search_charmap",
    ]
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row
    try:
        dump = {}
        for table in tables:
            rows = conn.execute(f"SELECT * FROM {table}").fetchall()
            dump[table] = sorted(tuple(row) for row in rows)
        return dump
    finally:
        conn.close()


def test_write_bundle_includes_phantom_group_with_zero_links(tmp_path):
    """A group created only to carry a link_shared_across_groups review flag
    (see test_link_shared_across_groups_conflict_flags_both_groups) never
    gets an actual link pointed at it -- store.recount() must still handle
    it and produce link_count == 0, needs_review == 1 rather than erroring
    or silently dropping the row."""
    same_link = "https://pan.quark.cn/s/fakeconflict2"
    row_a = _row(
        record_id=1, slug="s1", link=same_link,
        title="幽灵组剧 (2024)", media_title="幽灵组剧", remark="S01 4K WEB-DL",
    )
    row_b = _row(
        record_id=2, slug="s2", link=same_link,
        title="幽灵组剧 (2024)", media_title="幽灵组剧", remark="S02 4K WEB-DL", row_number=6,
    )
    agg = importer.aggregate(_rows_by_source(main=[row_a, row_b]))
    assert agg.totals["resource_groups"] == 2, "group B is created just to be flagged"

    path = importer.write_bundle(agg, tmp_path / "out", source_hashes={})
    conn = sqlite3.connect(str(path))
    try:
        rows = conn.execute(
            "SELECT link_count, needs_review FROM resource_group ORDER BY link_count"
        ).fetchall()
    finally:
        conn.close()
    assert rows == [(0, 1), (1, 1)]


def test_write_report_json_never_contains_link_or_access_code_values(tmp_path):
    agg = _small_agg()
    report_path = tmp_path / "report.json"
    sources_meta = [
        {"file": "HDHive official-group-backup.xlsx", "sheet": "影视分享", "kind": "main",
         "sha256": "deadbeef" * 8, "header_row": 4},
        {"file": "影巢_影视分享_整理.xlsx", "sheet": "影视分享", "kind": "supplement_video",
         "sha256": "cafebabe" * 8, "header_row": 4},
        {"file": "影巢_ED2K_整理.xlsx", "sheet": "磁力ED2K", "kind": "supplement_ed2k",
         "sha256": "0" * 64, "header_row": 4},
    ]
    importer.write_report(agg, sources_meta, report_path, "dry-run")
    text = report_path.read_text(encoding="utf-8")

    assert "http" not in text
    assert "ed2k://" not in text
    for code in ("ab12", "cd12", "ef34", "cd34", "cd56", "zz99"):
        assert code not in text

    mode = stat.S_IMODE(report_path.stat().st_mode)
    assert mode == 0o600

    report = json.loads(text)
    assert report["mode"] == "dry-run"
    assert report["status"] == "ok"
    assert report["totals"]["links"] == agg.totals["links"]
    # Real-data hotfix fields: present on every source entry and in
    # needs_review, zero here since _small_agg() has no bad timestamps or
    # shifted rows -- see the dedicated leak test below for the non-zero,
    # value-bearing case.
    for source in report["sources"]:
        assert source["rows_bad_timestamp"] == 0
        assert source["rows_shifted"] == 0
        # I7-I12/Minor counters: present on every source entry, zero here
        # since _small_agg() has none of these anomalies either -- see the
        # dedicated leak test below for the non-zero, value-bearing case.
        assert source["rows_error"] == 0
        assert source["rows_missing_link"] == 0
        assert source["rows_type_missing"] == 0
        assert source["rows_missing_record_id"] == 0
        assert source["rows_multi_link"] == 0
        assert source["links_from_multi_link_cells"] == 0
    assert report["needs_review"]["timestamp_unparsed"] == 0
    assert report["needs_review"]["row_shifted"] == 0
    assert report["conflicts"]["access_code_divergence"] == 0


def test_write_report_json_never_leaks_shifted_row_or_bad_timestamp_cell_values(tmp_path):
    """Extends the no-leak guarantee above to the two real-data hotfix
    fields: a shifted row's real link/access-code cells (living in
    different columns than usual) and a garbage timestamp string must still
    never reach the JSON report -- only the rows_shifted/rows_bad_timestamp
    counters and the row_shifted/timestamp_unparsed needs_review reasons
    may."""
    secret_url = "https://115cdn.com/s/swfakeleakshift1"
    secret_code = "zk77"
    bad_ts_marker = "not-a-real-date-zz99-fake-marker"

    shifted_row = {
        "类型": "影视", "记录ID": 1, "Slug/追更包Slug": "slug-leak-shift",
        "用户UID": 500001, "官组标记": "无能",
        "标题": "泄露检测剧 (2024)", "媒体标题": "泄露检测剧",
        "链接": None, "访问码": secret_url,
        "备注": secret_code, "创建时间": "S01 4K WEB-DL",
        "删除时间": datetime(2024, 1, 1), "_row_number": 5,
    }
    bad_ts_row = _row(
        record_id=2, slug="slug-leak-ts", link="https://pan.quark.cn/s/fakeleakts1",
        created=bad_ts_marker, row_number=6,
    )

    agg = importer.aggregate(_rows_by_source(main=[shifted_row, bad_ts_row]))
    report_path = tmp_path / "report.json"
    sources_meta = [
        {"file": "HDHive official-group-backup.xlsx", "sheet": "影视分享", "kind": "main",
         "sha256": "deadbeef" * 8, "header_row": 4},
    ]
    importer.write_report(agg, sources_meta, report_path, "dry-run")
    text = report_path.read_text(encoding="utf-8")

    assert secret_url not in text
    assert secret_code not in text
    assert bad_ts_marker not in text
    assert "http" not in text

    report = json.loads(text)
    assert report["sources"][0]["rows_shifted"] == 1
    assert report["sources"][0]["rows_bad_timestamp"] == 1
    assert report["needs_review"]["row_shifted"] == 1
    assert report["needs_review"]["timestamp_unparsed"] == 1


def test_write_report_json_never_leaks_contained_or_multi_link_cell_values(tmp_path):
    """Extends the no-leak guarantee to I7 (a row-error's cell value), I9
    (a blank-link row -- vacuously never leaks since nothing is stored, but
    checked anyway), and I11 (each fragment of a multi-link cell)."""
    bad_record_id_marker = "not-a-real-record-id-zz99-fake-marker"
    multi_ed2k_1 = "ed2k://|file|Fake.Leak.A|111111|00000000000000000000000000000011|/"
    multi_ed2k_2 = "ed2k://|file|Fake.Leak.B|222222|00000000000000000000000000000022|/"

    error_row = _row(record_id=bad_record_id_marker, slug="slug-leak-error", row_number=5)
    blank_link_row = _row(record_id=2, slug="slug-leak-blank", link="", row_number=6)
    multi_link_row = _row(
        record_id=3, slug="slug-leak-multi", link=multi_ed2k_1 + multi_ed2k_2, row_number=7,
    )

    agg = importer.aggregate(_rows_by_source(main=[error_row, blank_link_row, multi_link_row]))
    report_path = tmp_path / "report.json"
    sources_meta = [
        {"file": "HDHive official-group-backup.xlsx", "sheet": "影视分享", "kind": "main",
         "sha256": "deadbeef" * 8, "header_row": 4},
    ]
    importer.write_report(agg, sources_meta, report_path, "dry-run")
    text = report_path.read_text(encoding="utf-8")

    assert bad_record_id_marker not in text
    assert "ed2k://" not in text
    assert "Fake.Leak" not in text

    report = json.loads(text)
    assert report["sources"][0]["rows_error"] == 1
    assert report["sources"][0]["rows_missing_link"] == 1
    assert report["sources"][0]["rows_multi_link"] == 1
    assert report["sources"][0]["links_from_multi_link_cells"] == 2
    assert report["totals"]["links"] == 2


# ---------------------------------------------------------------------------
# End-to-end: library_sources fixture against EXPECTED
# ---------------------------------------------------------------------------


def _read_all_sources(source_dir: Path):
    rows_by_source = {}
    sources_meta = []
    for filename, sheet, kind in importer.SOURCES:
        header_row, rows = importer.read_sheet(source_dir / filename, sheet)
        rows_by_source[kind] = rows
        sources_meta.append({"file": filename, "sheet": sheet, "kind": kind, "header_row": header_row})
    cc_file, cc_sheet = importer.CROSS_CHECK
    _cc_header, cc_rows = importer.read_sheet(source_dir / cc_file, cc_sheet)
    rows_by_source["cross_check"] = cc_rows
    return rows_by_source, sources_meta


def test_full_pipeline_matches_expected_counts(library_sources):
    rows_by_source, _meta = _read_all_sources(library_sources)
    agg = importer.aggregate(rows_by_source)

    main_stat = agg.sources["main"]
    assert main_stat["rows_total"] == len(VIDEO_ROWS)
    assert main_stat["rows_video"] == len(VIDEO_ROWS)
    assert main_stat["rows_skipped_non_video"] == 0
    assert main_stat["rows_invalid_link"] == EXPECTED["video_non_url_rows"]
    assert main_stat["rows_deleted"] == EXPECTED["video_deleted_rows"]

    expected_unique_links = EXPECTED["video_unique_links"]
    assert (len(VIDEO_ROWS) - expected_unique_links) == agg.dedup["within_main_duplicate_links"]

    assert agg.dedup["supplement_video_matched"] == EXPECTED["supplement_video_duplicate_rows"]
    assert agg.dedup["supplement_video_added"] == EXPECTED["supplement_video_new_rows"]
    assert agg.dedup["supplement_ed2k_matched"] == EXPECTED["supplement_ed2k_duplicate_rows"]
    assert agg.dedup["supplement_ed2k_added"] == EXPECTED["supplement_ed2k_new_rows"]
    assert agg.dedup["cross_check_ed2k_sheet_subset"] is True

    expected_total_links = (
        expected_unique_links
        + EXPECTED["supplement_video_new_rows"]
        + EXPECTED["supplement_ed2k_new_rows"]
    )
    assert agg.totals["links"] == expected_total_links

    # provider distribution of the FINAL deduped link set, derived here from
    # the generator's own row lists (never hardcoded) via the same
    # parse_link() the importer uses, over the union of distinct
    # (provider, canonical) pairs across all three sources.
    from library_normalize import parse_link
    from collections import Counter

    expected_pairs = set()
    for row in VIDEO_ROWS + SUPPLEMENT_VIDEO_ROWS + SUPPLEMENT_ED2K_ROWS:
        info = parse_link(row["链接"], row["访问码"])
        expected_pairs.add((info.provider, info.canonical))
    expected_providers = Counter(provider for provider, _canonical in expected_pairs)
    assert dict(agg.providers) == dict(expected_providers)


def test_full_pipeline_media_type_distribution_covers_movie_tv_and_unknown(library_sources):
    """§4.6: infer_media_type()'s three outcomes must each be reachable
    through the real importer path (read_sheet -> aggregate), not just in
    isolated unit tests -- the fixture must carry at least one markerless
    title with a TV-ish cue (-> "unknown") alongside its movie/tv rows."""
    from collections import Counter

    rows_by_source, _meta = _read_all_sources(library_sources)
    agg = importer.aggregate(rows_by_source)

    counts = Counter(media["media_type"] for media in agg.media.values())
    assert counts["movie"] >= 1
    assert counts["tv"] >= 1
    assert counts["unknown"] >= 1


def test_full_pipeline_year_missing_matches_expected(library_sources):
    rows_by_source, _meta = _read_all_sources(library_sources)
    agg = importer.aggregate(rows_by_source)
    assert agg.year_missing == EXPECTED["video_no_year_rows"]


def test_full_pipeline_write_bundle_and_search_index(library_sources, tmp_path):
    rows_by_source, _meta = _read_all_sources(library_sources)
    agg = importer.aggregate(rows_by_source)
    path = importer.write_bundle(agg, tmp_path / "out", source_hashes={"x": "y"})

    store = ls.LibraryStore(path)
    stats = store.stats()
    assert stats["media_total"] == agg.totals["media"]
    assert stats["links_total"] == agg.totals["links"]
    assert store.meta_get("normalize_version") is not None
    assert store.meta_get("encrypted") == "0"

    conn = sqlite3.connect(str(path))
    try:
        doc_count = conn.execute("SELECT COUNT(*) FROM search_doc").fetchone()[0]
        pinyin_count = conn.execute(
            "SELECT COUNT(*) FROM search_doc WHERE pinyin_full IS NOT NULL"
        ).fetchone()[0]
    finally:
        conn.close()
    assert doc_count == agg.totals["media"]
    assert pinyin_count > 0, "pypinyin-backed pinyin_full should be populated by default"


def test_full_pipeline_group_display_titles_are_never_the_media_title(library_sources, tmp_path):
    rows_by_source, _meta = _read_all_sources(library_sources)
    agg = importer.aggregate(rows_by_source)
    path = importer.write_bundle(agg, tmp_path / "out", source_hashes={"x": "y"})

    conn = sqlite3.connect(str(path))
    try:
        rows = conn.execute(
            "SELECT resource_group.display_title, media.title_zh FROM resource_group "
            "JOIN media ON media.id = resource_group.media_id"
        ).fetchall()
    finally:
        conn.close()
    assert rows, "synthetic fixture should produce at least one resource group"
    for display_title, title_zh in rows:
        assert display_title != title_zh


def test_full_pipeline_report_never_contains_any_fixture_access_code(library_sources, tmp_path):
    """Scan the report produced from the *entire* synthetic fixture (not just
    a hand-picked subset) for every access code the fixture actually uses --
    both the "访问码" column and codes embedded in a link's query string --
    derived here from the generator's own row lists via ``parse_link``
    (never hand-typed), so a new scenario row can never silently go
    unchecked."""
    from library_normalize import parse_link

    all_codes = set()
    for row in VIDEO_ROWS + SUPPLEMENT_VIDEO_ROWS + SUPPLEMENT_ED2K_ROWS:
        info = parse_link(row["链接"], row["访问码"])
        if info.access_code:
            all_codes.add(info.access_code)
    assert all_codes, "fixture should exercise at least one access code"

    rows_by_source, _meta = _read_all_sources(library_sources)
    agg = importer.aggregate(rows_by_source)
    report_path = tmp_path / "full-report.json"
    importer.write_report(agg, [], report_path, "dry-run")
    text = report_path.read_text(encoding="utf-8")

    assert "http" not in text
    assert "ed2k://" not in text
    for code in all_codes:
        assert code not in text, f"access code leaked into report: {code!r}"


# ---------------------------------------------------------------------------
# main() CLI
# ---------------------------------------------------------------------------


def test_main_dry_run_writes_no_bundle(library_sources, tmp_path, capsys):
    out_dir = tmp_path / "out"
    report_path = tmp_path / "report.json"
    code = importer.main([
        "--source-dir", str(library_sources),
        "--output", str(out_dir),
        "--report", str(report_path),
    ])
    assert code == 0
    assert not out_dir.exists() or not any(out_dir.iterdir())
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["mode"] == "dry-run"


def test_main_write_creates_bundle_and_report(library_sources, tmp_path):
    out_dir = tmp_path / "out"
    report_path = tmp_path / "report.json"
    code = importer.main([
        "--source-dir", str(library_sources),
        "--output", str(out_dir),
        "--write",
        "--report", str(report_path),
    ])
    assert code == 0
    bundle = out_dir / "library-bundle.sqlite"
    assert bundle.is_file()
    assert stat.S_IMODE(bundle.stat().st_mode) == 0o600
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["mode"] == "write"
    assert report["dedup"]["cross_check_ed2k_sheet_subset"] is True


def test_main_write_is_idempotent(library_sources, tmp_path):
    out_dir = tmp_path / "out"
    argv = [
        "--source-dir", str(library_sources),
        "--output", str(out_dir),
        "--write",
        "--report", str(tmp_path / "report.json"),
    ]
    assert importer.main(argv) == 0
    dump_a = _dump_tables(out_dir / "library-bundle.sqlite")
    assert importer.main(argv) == 0
    dump_b = _dump_tables(out_dir / "library-bundle.sqlite")
    assert dump_a == dump_b


def _build_minimal_source_workbook(path: Path, sheets: dict) -> None:
    """Write a minimal .xlsx with a header row (importer.HEADER, row 1) and
    zero or more data rows, per sheet -- enough for read_sheet(), without
    depending on the full fixtures/library/make_workbooks.py generator."""
    from openpyxl import Workbook

    header = list(importer.HEADER)
    wb = Workbook()
    names = list(sheets.keys())
    first = wb.active
    first.title = names[0]
    first.append(header)
    for row in sheets[names[0]]:
        first.append(row)
    for name in names[1:]:
        ws = wb.create_sheet(name)
        ws.append(header)
        for row in sheets[name]:
            ws.append(row)
    wb.save(path)


def _minimal_video_row(record_id, slug, link):
    return [
        "影视", record_id, slug, 500001, "无能", "示例剧 (2024)", "示例剧",
        link, "", "S01 4K WEB-DL", datetime(2024, 1, 1), None,
    ]


def test_main_write_fails_and_writes_report_when_cross_check_finds_orphan_ed2k(tmp_path):
    """Integration test for the §4.8 rule 1 cross-check failure path: the
    official workbook's 磁力ED2K sheet contains an ED2K row absent from
    影视分享.  main(["--write", ...]) must exit 1, still write the report
    (with dedup.cross_check_ed2k_sheet_subset == false), and must NOT leave
    a bundle file behind (write_bundle is skipped once the cross-check is
    known to have failed, rather than writing then discarding it)."""
    main_link = "https://pan.quark.cn/s/fakecc01"
    ed2k_ok = "ed2k://|file|Fake.CC.OK|123456|0123456789ABCDEF0123456789ABCDEF01|/"
    ed2k_orphan = "ed2k://|file|Fake.CC.Orphan|999999|0123456789ABCDEF0123456789ABCDEFFF|/"

    source_dir = tmp_path / "source"
    source_dir.mkdir()
    _build_minimal_source_workbook(
        source_dir / "HDHive official-group-backup.xlsx",
        {
            "影视分享": [
                _minimal_video_row(1, "slug-main-1", main_link),
                _minimal_video_row(2, "slug-cc-ok", ed2k_ok),
            ],
            "磁力ED2K": [
                _minimal_video_row(2, "slug-cc-ok", ed2k_ok),
                _minimal_video_row(3, "slug-cc-orphan", ed2k_orphan),
            ],
        },
    )
    _build_minimal_source_workbook(source_dir / "影巢_影视分享_整理.xlsx", {"影视分享": []})
    _build_minimal_source_workbook(source_dir / "影巢_ED2K_整理.xlsx", {"磁力ED2K": []})

    out_dir = tmp_path / "out"
    report_path = tmp_path / "report.json"
    code = importer.main([
        "--source-dir", str(source_dir),
        "--output", str(out_dir),
        "--write",
        "--report", str(report_path),
    ])

    assert code == 1
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["dedup"]["cross_check_ed2k_sheet_subset"] is False
    assert not out_dir.exists() or not any(out_dir.iterdir()), "bundle must not be left behind"


def test_main_limit_truncates_rows(library_sources, tmp_path):
    report_path = tmp_path / "report.json"
    code = importer.main([
        "--source-dir", str(library_sources),
        "--output", str(tmp_path / "out"),
        "--report", str(report_path),
        "--limit", "5",
    ])
    assert code == 0
    report = json.loads(report_path.read_text(encoding="utf-8"))
    main_source = next(s for s in report["sources"] if s["file"] == "HDHive official-group-backup.xlsx")
    assert main_source["rows_total"] == 5


def test_main_missing_source_dir_argument_is_usage_error(tmp_path):
    # argparse itself raises SystemExit(2) for a missing required argument,
    # same as every other argparse-based CLI in this repo (see
    # scripts/enrich_media_tmdb.py) -- main() does not catch it.
    with pytest.raises(SystemExit) as excinfo:
        importer.main(["--output", str(tmp_path / "out")])
    assert excinfo.value.code == 2


def test_main_missing_source_file_is_input_error(tmp_path):
    empty_dir = tmp_path / "empty"
    empty_dir.mkdir()
    code = importer.main([
        "--source-dir", str(empty_dir),
        "--output", str(tmp_path / "out"),
    ])
    assert code == 1


def test_main_missing_source_file_writes_minimal_failure_report(tmp_path):
    empty_dir = tmp_path / "empty"
    empty_dir.mkdir()
    report_path = tmp_path / "report.json"
    code = importer.main([
        "--source-dir", str(empty_dir),
        "--output", str(tmp_path / "out"),
        "--report", str(report_path),
    ])
    assert code == 1
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report == {
        "version": 1,
        "generated_at": report["generated_at"],
        "mode": "dry-run",
        "status": "failed",
        "error": "SourceFileMissingError",
    }


def test_main_import_row_error_writes_minimal_failure_report(tmp_path):
    """A minimal, desensitised report -- no sheet/column names, which
    ImportRowError's own message would carry -- for a failure that happens
    before aggregate() ever runs."""
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    _build_minimal_source_workbook(source_dir / "影巢_影视分享_整理.xlsx", {"影视分享": []})
    _build_minimal_source_workbook(source_dir / "影巢_ED2K_整理.xlsx", {"磁力ED2K": []})

    from openpyxl import Workbook

    bad_header = [h for h in importer.HEADER if h != "访问码"]
    wb = Workbook()
    ws = wb.active
    ws.title = "影视分享"
    for col, name in enumerate(bad_header, start=1):
        ws.cell(row=1, column=col, value=name)
    ws.cell(row=2, column=1, value=1)
    wb.save(source_dir / "HDHive official-group-backup.xlsx")

    report_path = tmp_path / "report.json"
    code = importer.main([
        "--source-dir", str(source_dir),
        "--output", str(tmp_path / "out"),
        "--report", str(report_path),
    ])
    assert code == 1
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report == {
        "version": 1,
        "generated_at": report["generated_at"],
        "mode": "dry-run",
        "status": "failed",
        "error": "ImportRowError",
    }


def test_main_aggregate_import_row_error_writes_minimal_failure_report(tmp_path, monkeypatch):
    """Same minimal-report/exit-1 guarantee as
    test_main_import_row_error_writes_minimal_failure_report, but for an
    ImportRowError raised *during* aggregate() rather than during
    read_sheet() -- aggregate() runs after run()'s own read_sheet
    try/except block, so this failure must still be caught there, not
    escape main() uncaught.  Specifically exercises I8: a main-loop row
    failure (e.g. an unparseable 记录ID) is now contained per-row (I7) and
    no longer aborts the run, so this uses a cross-check-loop parse_link
    failure instead -- the one aggregate()-time failure that still
    propagates all the way out."""
    trigger = "https://pan.quark.cn/s/TRIGGER-CROSS-CHECK-BOOM-fake"
    real_parse_link = importer.parse_link

    def boom(raw_url, access_code_cell):
        if raw_url == trigger:
            raise ValueError("boom")
        return real_parse_link(raw_url, access_code_cell)

    monkeypatch.setattr(importer, "parse_link", boom)

    source_dir = tmp_path / "source"
    source_dir.mkdir()
    main_link = "https://pan.quark.cn/s/fakecrossboom1"
    _build_minimal_source_workbook(
        source_dir / "HDHive official-group-backup.xlsx",
        {
            "影视分享": [_minimal_video_row(1, "slug-main-1", main_link)],
            "磁力ED2K": [_minimal_video_row(1, "slug-main-1", trigger)],
        },
    )
    _build_minimal_source_workbook(source_dir / "影巢_影视分享_整理.xlsx", {"影视分享": []})
    _build_minimal_source_workbook(source_dir / "影巢_ED2K_整理.xlsx", {"磁力ED2K": []})

    report_path = tmp_path / "report.json"
    code = importer.main([
        "--source-dir", str(source_dir),
        "--output", str(tmp_path / "out"),
        "--report", str(report_path),
    ])
    assert code == 1
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report == {
        "version": 1,
        "generated_at": report["generated_at"],
        "mode": "dry-run",
        "status": "failed",
        "error": "ImportRowError",
    }


def test_main_require_hashes_fails_on_mismatch(library_sources, tmp_path):
    readme = library_sources / "README.md"
    readme.write_text(
        "## 输入完整性（SHA-256）\n\n"
        "| 文件 | SHA-256 |\n|---|---|\n"
        "| `HDHive official-group-backup.xlsx` | `" + ("0" * 64) + "` |\n"
        "| `影巢_影视分享_整理.xlsx` | `" + ("0" * 64) + "` |\n"
        "| `影巢_ED2K_整理.xlsx` | `" + ("0" * 64) + "` |\n",
        encoding="utf-8",
    )
    report_path = tmp_path / "report.json"
    code = importer.main([
        "--source-dir", str(library_sources),
        "--output", str(tmp_path / "out"),
        "--report", str(report_path),
        "--require-hashes",
    ])
    assert code == 1
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report == {
        "version": 1,
        "generated_at": report["generated_at"],
        "mode": "dry-run",
        "status": "failed",
        "error": "SourceHashMismatchError",
    }


def test_main_require_hashes_passes_on_match(library_sources, tmp_path):
    official = library_sources / "HDHive official-group-backup.xlsx"
    supp_video = library_sources / "影巢_影视分享_整理.xlsx"
    supp_ed2k = library_sources / "影巢_ED2K_整理.xlsx"
    real_hash = importer._sha256_file(official)
    real_hash_video = importer._sha256_file(supp_video)
    real_hash_ed2k = importer._sha256_file(supp_ed2k)
    readme = library_sources / "README.md"
    readme.write_text(
        "## 输入完整性（SHA-256）\n\n"
        "| 文件 | SHA-256 |\n|---|---|\n"
        f"| `HDHive official-group-backup.xlsx` | `{real_hash}` |\n"
        f"| `影巢_影视分享_整理.xlsx` | `{real_hash_video}` |\n"
        f"| `影巢_ED2K_整理.xlsx` | `{real_hash_ed2k}` |\n",
        encoding="utf-8",
    )
    code = importer.main([
        "--source-dir", str(library_sources),
        "--output", str(tmp_path / "out"),
        "--report", str(tmp_path / "report.json"),
        "--require-hashes",
    ])
    assert code == 0


def test_write_bundle_points_sqlite_temp_files_at_output_dir(tmp_path, monkeypatch):
    """The importer's VACUUM/index temp files must land next to the bundle,
    never in a (possibly tiny tmpfs) /tmp -- mirrors install_bundle."""
    import os as _os
    import sqlite3 as _sqlite3

    monkeypatch.delenv("SQLITE_TMPDIR", raising=False)
    out_dir = tmp_path / "bundle-out"
    agg = importer.aggregate({"main": VIDEO_ROWS})
    importer.write_bundle(agg, out_dir, source_hashes={})
    probe = _sqlite3.connect(":memory:")
    try:
        assert probe.execute("PRAGMA temp_store_directory").fetchone()[0] == str(out_dir)
    finally:
        probe.close()
    assert _os.environ["SQLITE_TMPDIR"] == str(out_dir)
